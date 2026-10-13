"""
منشئ استمارة تقييم أداء التدريسيين بصيغة Microsoft Word (.docx)
مطابق تماماً لاستمارة رقم (21) للعام الدراسي 2025-2026 - وزارة التعليم العالي والبحث العلمي
مع تنسيق الجداول الرسمية والخلفيات الصفراء ومحاذاة RTL وملحق الوثائق والمرفقات الثبوتية.
"""
import os
from typing import Dict, Any, List
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

from core.scoring_engine import calculate_evaluation, calculate_evaluation_form39, number_to_arabic_words

YELLOW_HEX = "FFF200"
GRAY_HEX = "E0E0E0"
LIGHT_YELLOW = "FFFFCC"


def set_cell_background(cell, fill_hex: str):
    """تعيين لون خلفية خلية الجدول"""
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)


def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """تعيين هوامش خلية الجدول"""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'''<w:tcMar {nsdecls("w")}>
        <w:top w:w="{top}" w:type="dxa"/>
        <w:bottom w:w="{bottom}" w:type="dxa"/>
        <w:left w:w="{left}" w:type="dxa"/>
        <w:right w:w="{right}" w:type="dxa"/>
    </w:tcMar>''')
    tcPr.append(tcMar)


def set_table_rtl(table):
    """جعل الجدول يدعم الاتجاه من اليمين إلى اليسار (RTL) للغة العربية"""
    tblPr = table._tbl.tblPr
    bidiVisual = parse_xml(f'<w:bidiVisual {nsdecls("w")}/>')
    tblPr.append(bidiVisual)


def set_table_borders(table, color="000000", sz="4", val="single"):
    """تعيين حدود خلايا الجدول بالكامل"""
    tblPr = table._tbl.tblPr
    borders = parse_xml(f'''
        <w:tblBorders {nsdecls("w")}>
            <w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:left w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:right w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:insideV w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
        </w:tblBorders>
    ''')
    tblPr.append(borders)


def format_paragraph(p, text="", bold=False, italic=False, size=11, align=WD_ALIGN_PARAGRAPH.RIGHT, color=None):
    """تنسيق الفقرة والخط العربي"""
    p.alignment = align
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.15
    # تفعيل RTL للفقرة
    pPr = p._p.get_or_add_pPr()
    bidi = parse_xml(f'<w:bidi {nsdecls("w")} w:val="1"/>')
    pPr.append(bidi)
    
    if text:
        run = p.add_run(text)
        run.bold = bold
        run.italic = italic
        run.font.name = "Arial"
        run.font.size = Pt(size)
        if color:
            run.font.color.rgb = color
        # تعيين لغة الخط العربي
        rPr = run._r.get_or_add_rPr()
        rFonts = parse_xml(f'<w:rFonts {nsdecls("w")} w:ascii="Arial" w:hAnsi="Arial" w:cs="Arial"/>')
        rPr.append(rFonts)
    return p


def create_form_21_docx(data: Dict[str, Any], output_path: str, attachments: List[Dict[str, Any]] = None) -> str:
    """
    إنشاء ملف Word كامل للاستمارة 21 مع ملحق المرفقات
    """
    eval_res = calculate_evaluation(data)
    personal = data.get("personal_info", {})
    attachments = attachments or []

    doc = Document()

    # ضبط هوامش الصفحة
    for section in doc.sections:
        section.top_margin = Inches(0.5)
        section.bottom_margin = Inches(0.5)
        section.left_margin = Inches(0.5)
        section.right_margin = Inches(0.5)

    # =========================================================================
    # الصفحة الأولى : الترويسة والبيانات الرئيسة
    # =========================================================================
    header_table = doc.add_table(rows=1, cols=2)
    set_table_rtl(header_table)
    header_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    header_table.autofit = True

    # يمين: الترويسة الوزارية
    c_right = header_table.cell(0, 0)
    p_r = c_right.paragraphs[0]
    format_paragraph(p_r, "جمهورية العراق", bold=True, size=11, align=WD_ALIGN_PARAGRAPH.CENTER)
    p_r2 = c_right.add_paragraph()
    format_paragraph(p_r2, "وزارة التعليم العالي والبحث العلمي\nجهاز الإشراف والتقويم العلمي\nدائرة ضمان الجودة والاعتماد الأكاديمي\nقسم تقويم الأداء المؤسسي", bold=True, size=10, align=WD_ALIGN_PARAGRAPH.CENTER)

    # يسار: رقم وترميز الاستمارة
    c_left = header_table.cell(0, 1)
    p_l = c_left.paragraphs[0]
    form_no = personal.get("form_no", "2025/....")
    form_code = personal.get("form_code", "AGY-....")
    format_paragraph(p_l, f"رقم الاستمارة: {form_no}\nترميز الاستمارة: {form_code}", bold=True, size=11, align=WD_ALIGN_PARAGRAPH.LEFT)

    # عنوان الاستمارة الإطاري
    title_box = doc.add_table(rows=1, cols=1)
    set_table_rtl(title_box)
    set_table_borders(title_box, sz="8")
    c_title = title_box.cell(0, 0)
    p_t = c_title.paragraphs[0]
    format_paragraph(p_t, "استمارة رقم (21): تقييم أداء أعضاء الهيئة التدريسية للعام الدراسي 2025-2026", bold=True, size=13, align=WD_ALIGN_PARAGRAPH.CENTER)
    set_cell_background(c_title, LIGHT_YELLOW)

    p_space = doc.add_paragraph()
    p_space.paragraph_format.space_before = Pt(6)

    # الجامعة والكلية والقسم
    p_org = doc.add_paragraph()
    format_paragraph(p_org, f"البيانات الرئيسة :\nالجامعة: {personal.get('university', '......................')}    الكلية: {personal.get('college', '......................')}    القسم / الفرع: {personal.get('department', '......................')}", bold=True, size=11)

    # جدول البيانات الشخصية (مطابق للصفحة 1)
    p1_table = doc.add_table(rows=16, cols=5)
    set_table_rtl(p1_table)
    set_table_borders(p1_table)
    p1_table.alignment = WD_TABLE_ALIGNMENT.CENTER

    # Row 0: الاسم الرباعي واللقب (عناوين صفراء)
    headers_r0 = ["اللقب", "اسم جد الأب", "اسم الجد", "اسم الأب", "الاسم"]
    for i, h in enumerate(headers_r0):
        c = p1_table.cell(0, i)
        set_cell_background(c, YELLOW_HEX)
        format_paragraph(c.paragraphs[0], h, bold=True, size=10, align=WD_ALIGN_PARAGRAPH.CENTER)

    # Row 1: قيم الاسم الرباعي
    vals_r1 = [
        personal.get("last_name", ""),
        personal.get("great_grandfather_name", ""),
        personal.get("grandfather_name", ""),
        personal.get("father_name", ""),
        personal.get("first_name", "")
    ]
    for i, v in enumerate(vals_r1):
        c = p1_table.cell(1, i)
        format_paragraph(c.paragraphs[0], v, bold=False, size=10, align=WD_ALIGN_PARAGRAPH.CENTER)

    # Row 2: اسم الأم الثلاثي
    p1_table.cell(2, 0).merge(p1_table.cell(2, 1))
    p1_table.cell(2, 2).merge(p1_table.cell(2, 3))
    # الآن لدينا 3 أعمدة
    set_cell_background(p1_table.cell(2, 0), YELLOW_HEX)
    format_paragraph(p1_table.cell(2, 0).paragraphs[0], "اسم جد الأم", bold=True, size=10, align=WD_ALIGN_PARAGRAPH.CENTER)
    set_cell_background(p1_table.cell(2, 2), YELLOW_HEX)
    format_paragraph(p1_table.cell(2, 2).paragraphs[0], "اسم والد الأم", bold=True, size=10, align=WD_ALIGN_PARAGRAPH.CENTER)
    set_cell_background(p1_table.cell(2, 4), YELLOW_HEX)
    format_paragraph(p1_table.cell(2, 4).paragraphs[0], "اسم الأم", bold=True, size=10, align=WD_ALIGN_PARAGRAPH.CENTER)

    # Row 3: قيم اسم الأم
    p1_table.cell(3, 0).merge(p1_table.cell(3, 1))
    p1_table.cell(3, 2).merge(p1_table.cell(3, 3))
    format_paragraph(p1_table.cell(3, 0).paragraphs[0], personal.get("mother_grandfather_name", ""), size=10, align=WD_ALIGN_PARAGRAPH.CENTER)
    format_paragraph(p1_table.cell(3, 2).paragraphs[0], personal.get("mother_father_name", ""), size=10, align=WD_ALIGN_PARAGRAPH.CENTER)
    format_paragraph(p1_table.cell(3, 4).paragraphs[0], personal.get("mother_name", ""), size=10, align=WD_ALIGN_PARAGRAPH.CENTER)

    # Row 4: البطاقة الموحدة / الجنسية
    p1_table.cell(4, 2).merge(p1_table.cell(4, 4))
    set_cell_background(p1_table.cell(4, 0), YELLOW_HEX)
    format_paragraph(p1_table.cell(4, 0).paragraphs[0], "الصحيفة", bold=True, size=10, align=WD_ALIGN_PARAGRAPH.CENTER)
    set_cell_background(p1_table.cell(4, 1), YELLOW_HEX)
    format_paragraph(p1_table.cell(4, 1).paragraphs[0], "السجل", bold=True, size=10, align=WD_ALIGN_PARAGRAPH.CENTER)
    set_cell_background(p1_table.cell(4, 2), YELLOW_HEX)
    format_paragraph(p1_table.cell(4, 2).paragraphs[0], "رقم الجنسية أو البطاقة الموحدة", bold=True, size=10, align=WD_ALIGN_PARAGRAPH.CENTER)

    # Row 5: قيم الجنسية
    p1_table.cell(5, 2).merge(p1_table.cell(5, 4))
    format_paragraph(p1_table.cell(5, 0).paragraphs[0], personal.get("page_no", ""), size=10, align=WD_ALIGN_PARAGRAPH.CENTER)
    format_paragraph(p1_table.cell(5, 1).paragraphs[0], personal.get("registry_no", ""), size=10, align=WD_ALIGN_PARAGRAPH.CENTER)
    format_paragraph(p1_table.cell(5, 2).paragraphs[0], personal.get("national_id", ""), size=10, align=WD_ALIGN_PARAGRAPH.CENTER)

    # Row 6: تاريخ الإصدار
    p1_table.cell(6, 0).merge(p1_table.cell(6, 1))
    p1_table.cell(6, 3).merge(p1_table.cell(6, 4))
    set_cell_background(p1_table.cell(6, 0), YELLOW_HEX)
    format_paragraph(p1_table.cell(6, 0).paragraphs[0], "يوم الإصدار", bold=True, size=10, align=WD_ALIGN_PARAGRAPH.CENTER)
    set_cell_background(p1_table.cell(6, 2), YELLOW_HEX)
    format_paragraph(p1_table.cell(6, 2).paragraphs[0], "شهر الإصدار", bold=True, size=10, align=WD_ALIGN_PARAGRAPH.CENTER)
    set_cell_background(p1_table.cell(6, 3), YELLOW_HEX)
    format_paragraph(p1_table.cell(6, 3).paragraphs[0], "سنة الإصدار", bold=True, size=10, align=WD_ALIGN_PARAGRAPH.CENTER)

    # Row 7: قيم تاريخ الإصدار
    p1_table.cell(7, 0).merge(p1_table.cell(7, 1))
    p1_table.cell(7, 3).merge(p1_table.cell(7, 4))
    format_paragraph(p1_table.cell(7, 0).paragraphs[0], personal.get("issue_day", ""), size=10, align=WD_ALIGN_PARAGRAPH.CENTER)
    format_paragraph(p1_table.cell(7, 2).paragraphs[0], personal.get("issue_month", ""), size=10, align=WD_ALIGN_PARAGRAPH.CENTER)
    format_paragraph(p1_table.cell(7, 3).paragraphs[0], personal.get("issue_year", ""), size=10, align=WD_ALIGN_PARAGRAPH.CENTER)

    # Row 8: الشهادة والتعيين
    p1_table.cell(8, 2).merge(p1_table.cell(8, 3))
    set_cell_background(p1_table.cell(8, 0), YELLOW_HEX)
    format_paragraph(p1_table.cell(8, 0).paragraphs[0], "يوم منح الشهادة", bold=True, size=9, align=WD_ALIGN_PARAGRAPH.CENTER)
    set_cell_background(p1_table.cell(8, 1), YELLOW_HEX)
    format_paragraph(p1_table.cell(8, 1).paragraphs[0], "شهر منح الشهادة", bold=True, size=9, align=WD_ALIGN_PARAGRAPH.CENTER)
    set_cell_background(p1_table.cell(8, 2), YELLOW_HEX)
    format_paragraph(p1_table.cell(8, 2).paragraphs[0], "رقم وتاريخ الأمر الوزاري أو الجامعي", bold=True, size=9, align=WD_ALIGN_PARAGRAPH.CENTER)
    set_cell_background(p1_table.cell(8, 4), YELLOW_HEX)
    format_paragraph(p1_table.cell(8, 4).paragraphs[0], "الشهادة", bold=True, size=9, align=WD_ALIGN_PARAGRAPH.CENTER)

    # Row 9: قيم الشهادة
    p1_table.cell(9, 2).merge(p1_table.cell(9, 3))
    format_paragraph(p1_table.cell(9, 0).paragraphs[0], personal.get("degree_day", ""), size=10, align=WD_ALIGN_PARAGRAPH.CENTER)
    format_paragraph(p1_table.cell(9, 1).paragraphs[0], personal.get("degree_month", ""), size=10, align=WD_ALIGN_PARAGRAPH.CENTER)
    format_paragraph(p1_table.cell(9, 2).paragraphs[0], personal.get("order_no_and_date", ""), size=10, align=WD_ALIGN_PARAGRAPH.CENTER)
    format_paragraph(p1_table.cell(9, 4).paragraphs[0], personal.get("degree", ""), size=10, align=WD_ALIGN_PARAGRAPH.CENTER)

    # Row 10: جهة منح الشهادة
    p1_table.cell(10, 0).merge(p1_table.cell(10, 1))
    set_cell_background(p1_table.cell(10, 0), YELLOW_HEX)
    format_paragraph(p1_table.cell(10, 0).paragraphs[0], "القسم", bold=True, size=10, align=WD_ALIGN_PARAGRAPH.CENTER)
    set_cell_background(p1_table.cell(10, 2), YELLOW_HEX)
    format_paragraph(p1_table.cell(10, 2).paragraphs[0], "الكلية", bold=True, size=10, align=WD_ALIGN_PARAGRAPH.CENTER)
    set_cell_background(p1_table.cell(10, 3), YELLOW_HEX)
    format_paragraph(p1_table.cell(10, 3).paragraphs[0], "الجامعة", bold=True, size=10, align=WD_ALIGN_PARAGRAPH.CENTER)
    set_cell_background(p1_table.cell(10, 4), YELLOW_HEX)
    format_paragraph(p1_table.cell(10, 4).paragraphs[0], "البلد المانح", bold=True, size=10, align=WD_ALIGN_PARAGRAPH.CENTER)

    # Row 11: قيم جهة منح الشهادة
    p1_table.cell(11, 0).merge(p1_table.cell(11, 1))
    format_paragraph(p1_table.cell(11, 0).paragraphs[0], personal.get("granting_dept", ""), size=10, align=WD_ALIGN_PARAGRAPH.CENTER)
    format_paragraph(p1_table.cell(11, 2).paragraphs[0], personal.get("granting_college", ""), size=10, align=WD_ALIGN_PARAGRAPH.CENTER)
    format_paragraph(p1_table.cell(11, 3).paragraphs[0], personal.get("granting_univ", ""), size=10, align=WD_ALIGN_PARAGRAPH.CENTER)
    format_paragraph(p1_table.cell(11, 4).paragraphs[0], personal.get("granting_country", ""), size=10, align=WD_ALIGN_PARAGRAPH.CENTER)

    # Row 12: الاختصاص
    p1_table.cell(12, 0).merge(p1_table.cell(12, 2))
    p1_table.cell(12, 3).merge(p1_table.cell(12, 4))
    set_cell_background(p1_table.cell(12, 0), YELLOW_HEX)
    format_paragraph(p1_table.cell(12, 0).paragraphs[0], "الاختصاص الدقيق", bold=True, size=10, align=WD_ALIGN_PARAGRAPH.CENTER)
    set_cell_background(p1_table.cell(12, 3), YELLOW_HEX)
    format_paragraph(p1_table.cell(12, 3).paragraphs[0], "الاختصاص العام", bold=True, size=10, align=WD_ALIGN_PARAGRAPH.CENTER)

    # Row 13: قيم الاختصاص
    p1_table.cell(13, 0).merge(p1_table.cell(13, 2))
    p1_table.cell(13, 3).merge(p1_table.cell(13, 4))
    format_paragraph(p1_table.cell(13, 0).paragraphs[0], personal.get("specific_specialty", ""), size=10, align=WD_ALIGN_PARAGRAPH.CENTER)
    format_paragraph(p1_table.cell(13, 3).paragraphs[0], personal.get("general_specialty", ""), size=10, align=WD_ALIGN_PARAGRAPH.CENTER)

    # Row 14: اللقب العلمي
    set_cell_background(p1_table.cell(14, 0), YELLOW_HEX)
    format_paragraph(p1_table.cell(14, 0).paragraphs[0], "اليوم", bold=True, size=10, align=WD_ALIGN_PARAGRAPH.CENTER)
    set_cell_background(p1_table.cell(14, 1), YELLOW_HEX)
    format_paragraph(p1_table.cell(14, 1).paragraphs[0], "الشهر", bold=True, size=10, align=WD_ALIGN_PARAGRAPH.CENTER)
    set_cell_background(p1_table.cell(14, 2), YELLOW_HEX)
    format_paragraph(p1_table.cell(14, 2).paragraphs[0], "السنة", bold=True, size=10, align=WD_ALIGN_PARAGRAPH.CENTER)
    set_cell_background(p1_table.cell(14, 3), YELLOW_HEX)
    format_paragraph(p1_table.cell(14, 3).paragraphs[0], "الجهة المانحة", bold=True, size=10, align=WD_ALIGN_PARAGRAPH.CENTER)
    set_cell_background(p1_table.cell(14, 4), YELLOW_HEX)
    format_paragraph(p1_table.cell(14, 4).paragraphs[0], "اللقب العلمي", bold=True, size=10, align=WD_ALIGN_PARAGRAPH.CENTER)

    # Row 15: قيم اللقب العلمي
    format_paragraph(p1_table.cell(15, 0).paragraphs[0], personal.get("title_day", ""), size=10, align=WD_ALIGN_PARAGRAPH.CENTER)
    format_paragraph(p1_table.cell(15, 1).paragraphs[0], personal.get("title_month", ""), size=10, align=WD_ALIGN_PARAGRAPH.CENTER)
    format_paragraph(p1_table.cell(15, 2).paragraphs[0], personal.get("title_year", ""), size=10, align=WD_ALIGN_PARAGRAPH.CENTER)
    format_paragraph(p1_table.cell(15, 3).paragraphs[0], personal.get("title_granter", ""), size=10, align=WD_ALIGN_PARAGRAPH.CENTER)
    format_paragraph(p1_table.cell(15, 4).paragraphs[0], personal.get("academic_title", ""), size=10, align=WD_ALIGN_PARAGRAPH.CENTER)

    # جدول الهاتف والبريد
    p1_sub_table = doc.add_table(rows=2, cols=2)
    set_table_rtl(p1_sub_table)
    set_table_borders(p1_sub_table)
    p1_sub_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_cell_background(p1_sub_table.cell(0, 0), YELLOW_HEX)
    set_cell_background(p1_sub_table.cell(0, 1), YELLOW_HEX)
    format_paragraph(p1_sub_table.cell(0, 0).paragraphs[0], "البريد الالكتروني", bold=True, size=10, align=WD_ALIGN_PARAGRAPH.CENTER)
    format_paragraph(p1_sub_table.cell(0, 1).paragraphs[0], "رقم الموبايل", bold=True, size=10, align=WD_ALIGN_PARAGRAPH.CENTER)
    format_paragraph(p1_sub_table.cell(1, 0).paragraphs[0], personal.get("email", ""), size=10, align=WD_ALIGN_PARAGRAPH.CENTER)
    format_paragraph(p1_sub_table.cell(1, 1).paragraphs[0], personal.get("phone", ""), size=10, align=WD_ALIGN_PARAGRAPH.CENTER)

    p_note = doc.add_paragraph()
    format_paragraph(p_note, "ملاحظة: يتحمل صاحب الاستمارة مسؤولية صحة المعلومات الواردة في الاستمارة.", bold=True, size=10, align=WD_ALIGN_PARAGRAPH.RIGHT)
    format_paragraph(doc.add_paragraph(), "5 - 1", bold=True, size=9, align=WD_ALIGN_PARAGRAPH.CENTER)

    doc.add_page_break()

    # =========================================================================
    # الصفحة الثانية : المحور الأول والمحور الثاني
    # =========================================================================
    p_ax1_head = doc.add_paragraph()
    format_paragraph(p_ax1_head, "المحور الأول : جودة التدريس والتعليم والالتزام الوظيفي يملئ من قبل اللجنة العلمية بعد ان يقدم صاحب العلاقة ما يؤيد\nالفقرات ( 1 ، 3 ، 4)، تملئ الفقرتين ( 2 ، 5) من قبل المسؤول المباشر ( %50) (التدريسي غير المكلف بمهام تدريسية تحتسب له الفقرة (5) فقط وتمثل درجة المحور باكمله)", bold=True, size=10)

    ax1_table = doc.add_table(rows=7, cols=5)
    set_table_rtl(ax1_table)
    set_table_borders(ax1_table)
    ax1_table.alignment = WD_TABLE_ALIGNMENT.CENTER

    th_ax1 = ["ت", "الفقرات", "التوصيف", "تفصيل الدرجة", "الدرجة المعطاة"]
    for i, h in enumerate(th_ax1):
        c = ax1_table.cell(0, i)
        set_cell_background(c, GRAY_HEX)
        format_paragraph(c.paragraphs[0], h, bold=True, size=10, align=WD_ALIGN_PARAGRAPH.CENTER)

    ax1_items = eval_res["axis1"]["items"]
    rows_ax1_data = [
        ("1", "المقررات التي قام بتدريسها", "تمنح (20) درجة لكل مقرر دراسي فصلي او سنوي (اولية او عليا) (نظري او عملي). الدرجة القصوى (20)", str(ax1_items["p1"]), f"{ax1_items['p1']}"),
        ("2", "ادارة الصف والعلاقة مع الطلبة واثارة دافعيتهم", "التقييم من خلال استمارة استبيان توزع على الطلبة (وفق الاستبانة المرفقة). الدرجة القصوى (20)", str(ax1_items["p2"]), f"{ax1_items['p2']}"),
        ("3", "التعليم المدمج", "تمنح (5) درجات لكل فقرة (استخدام طرائق متعددة، أمثلة تطبيقية، وسائل إيضاح، نشر محاضرات). الدرجة القصوى (20)", str(ax1_items["p3"]), f"{ax1_items['p3']}"),
        ("4", "وصف المقرر الدراسي وتحديثه والاساليب المستعملة", "تمنح (4) درجات لكل فقرة (وصف المقرر، مفردات المقرر، مقترحات التطوير، أساليب تقييم متنوعة، تغذية راجعة). القصوى (20)", str(ax1_items["p4"]), f"{ax1_items['p4']}"),
        ("5", "الالتزام الوظيفي", "تمنح (4) درجات لكل فقرة (20 درجة لكل فقرة للتدريسي غير المكلف). الدرجة القصوى (20) أو (100 لغير المكلف)", str(ax1_items["p5"]), f"{ax1_items['p5']}")
    ]

    for idx, (t, fq, desc, det, scr) in enumerate(rows_ax1_data, start=1):
        row = ax1_table.rows[idx]
        format_paragraph(row.cells[0].paragraphs[0], t, size=9, align=WD_ALIGN_PARAGRAPH.CENTER)
        format_paragraph(row.cells[1].paragraphs[0], fq, bold=True, size=9)
        format_paragraph(row.cells[2].paragraphs[0], desc, size=8)
        format_paragraph(row.cells[3].paragraphs[0], det, size=9, align=WD_ALIGN_PARAGRAPH.CENTER)
        format_paragraph(row.cells[4].paragraphs[0], scr, bold=True, size=10, align=WD_ALIGN_PARAGRAPH.CENTER)

    # المجموع المحور 1
    total_row_ax1 = ax1_table.rows[6]
    total_row_ax1.cells[0].merge(total_row_ax1.cells[3])
    format_paragraph(total_row_ax1.cells[0].paragraphs[0], "الدرجة النهائية للمحور الأول (القصوى 100)", bold=True, size=10, align=WD_ALIGN_PARAGRAPH.CENTER)
    format_paragraph(total_row_ax1.cells[4].paragraphs[0], f"{eval_res['axis1']['raw_total']}", bold=True, size=11, align=WD_ALIGN_PARAGRAPH.CENTER)

    # بداية المحور الثاني
    p_ax2_head = doc.add_paragraph()
    p_ax2_head.paragraph_format.space_before = Pt(8)
    format_paragraph(p_ax2_head, "المحور الثاني: النشاط العلمي والبحثي %30 يملئ من قبل اللجنة العلمية بعد ان تقدم الوثائق من قبل صاحب العلاقة المشمول بالتقييم", bold=True, size=10)

    ax2_table = doc.add_table(rows=5, cols=5)
    set_table_rtl(ax2_table)
    set_table_borders(ax2_table)
    ax2_table.alignment = WD_TABLE_ALIGNMENT.CENTER

    for i, h in enumerate(th_ax1):
        c = ax2_table.cell(0, i)
        set_cell_background(c, GRAY_HEX)
        format_paragraph(c.paragraphs[0], h, bold=True, size=10, align=WD_ALIGN_PARAGRAPH.CENTER)

    ax2_items = eval_res["axis2"]["items"]
    rows_ax2_data = [
        ("1", "البحوث العلمية المنشورة في المجلات المفهرسة ضمن المستوعبات العالمية", "تمنح (60) درجة لبحوث المستوعبات العالمية (Clarivate / Scopus) بمعامل CiteScore>=1، وتمنح (30) درجة للبحوث بمعامل CiteScore<1. الدرجة القصوى (60)", str(ax2_items["p1"]), f"{ax2_items['p1']}"),
        ("2", "البحوث العلمية المنشورة والكتب المؤلفة", "المجلات العربية والمحلية والمؤتمرات الدولية (25 درجة) والمحلية (15 درجة)، والكتب المؤلفة والمترجمة (20، 15، 10 درجات). القصوى (25)", str(ax2_items["p2"]), f"{ax2_items['p2']}"),
        ("3", "الاشراف على الطلبة والنشاطات العلمية الاخرى", "إشراف دكتوراه (15/10)، ماجستير (12/7)، دبلوم عالي (9/5)، مشاريع تخرج (6)، تقويم علمي (5). القصوى (15)", str(ax2_items["p3"]), f"{ax2_items['p3']}")
    ]

    for idx, (t, fq, desc, det, scr) in enumerate(rows_ax2_data, start=1):
        row = ax2_table.rows[idx]
        format_paragraph(row.cells[0].paragraphs[0], t, size=9, align=WD_ALIGN_PARAGRAPH.CENTER)
        format_paragraph(row.cells[1].paragraphs[0], fq, bold=True, size=9)
        format_paragraph(row.cells[2].paragraphs[0], desc, size=8)
        format_paragraph(row.cells[3].paragraphs[0], det, size=9, align=WD_ALIGN_PARAGRAPH.CENTER)
        format_paragraph(row.cells[4].paragraphs[0], scr, bold=True, size=10, align=WD_ALIGN_PARAGRAPH.CENTER)

    total_row_ax2 = ax2_table.rows[4]
    total_row_ax2.cells[0].merge(total_row_ax2.cells[3])
    format_paragraph(total_row_ax2.cells[0].paragraphs[0], "الدرجة النهائية للمحور الثاني (القصوى 100)", bold=True, size=10, align=WD_ALIGN_PARAGRAPH.CENTER)
    format_paragraph(total_row_ax2.cells[4].paragraphs[0], f"{eval_res['axis2']['raw_total']}", bold=True, size=11, align=WD_ALIGN_PARAGRAPH.CENTER)

    format_paragraph(doc.add_paragraph(), "5 - 2", bold=True, size=9, align=WD_ALIGN_PARAGRAPH.CENTER)
    doc.add_page_break()

    # =========================================================================
    # الصفحة الثالثة : المحور الثالث (الجانب التربوي والتعليم المستمر)
    # =========================================================================
    p_ax3_head = doc.add_paragraph()
    format_paragraph(p_ax3_head, "المحور الثالث: الجانب التربوي والارشادي والتعليم المستمر 20 %", bold=True, size=11)

    ax3_table = doc.add_table(rows=6, cols=5)
    set_table_rtl(ax3_table)
    set_table_borders(ax3_table)
    ax3_table.alignment = WD_TABLE_ALIGNMENT.CENTER

    for i, h in enumerate(th_ax1):
        c = ax3_table.cell(0, i)
        set_cell_background(c, GRAY_HEX)
        format_paragraph(c.paragraphs[0], h, bold=True, size=10, align=WD_ALIGN_PARAGRAPH.CENTER)

    ax3_items = eval_res["axis3"]["items"]
    rows_ax3_data = [
        ("1", "المشاركة في اللجان الدائمية والمؤقته داخل وزارة التعليم العالي", "عضو لجنة امتحانية (30)، لجنة دائمية (20)، لجنة مؤقتة (10)، امتحان تقويمي (5)، لجان سمنار (3). القصوى (30) درجة", str(ax3_items["p1"]), f"{ax3_items['p1']}"),
        ("2", "المشاركة في لجان التعليم المستمر والجودة", "محاضر في التعليم المستمر (10)، حضور (6 - حضورين فقط)، دورات طرائق التدريس الحديثة (8). القصوى (20) درجة", str(ax3_items["p2"]), f"{ax3_items['p2']}"),
        ("3", "كتب الشكر والتقدير او الشهادة التقديرية خلال عام التقييم", "من الوزير (20 شكر/10 تثمين)، وكيل وزير أو رئيس جامعة (15/8)، عميد أو مساعد رئيس (10/3). القصوى (20) درجة", str(ax3_items["p3"]), f"{ax3_items['p3']}"),
        ("4", "الزيارات الميدانية والحقلية والاعمال التطوعية والمجتمعية", "زيارة ميدانية للإشراف (12)، عمل تطوعي بالوزارة/الجامعة (8)، نشاطات وخدمة مجتمع خارج الوزارة (6). القصوى (30) درجة", str(ax3_items["p4"]), f"{ax3_items['p4']}")
    ]

    for idx, (t, fq, desc, det, scr) in enumerate(rows_ax3_data, start=1):
        row = ax3_table.rows[idx]
        format_paragraph(row.cells[0].paragraphs[0], t, size=9, align=WD_ALIGN_PARAGRAPH.CENTER)
        format_paragraph(row.cells[1].paragraphs[0], fq, bold=True, size=9)
        format_paragraph(row.cells[2].paragraphs[0], desc, size=8)
        format_paragraph(row.cells[3].paragraphs[0], det, size=9, align=WD_ALIGN_PARAGRAPH.CENTER)
        format_paragraph(row.cells[4].paragraphs[0], scr, bold=True, size=10, align=WD_ALIGN_PARAGRAPH.CENTER)

    total_row_ax3 = ax3_table.rows[5]
    total_row_ax3.cells[0].merge(total_row_ax3.cells[3])
    format_paragraph(total_row_ax3.cells[0].paragraphs[0], "الدرجة النهائية للمحور الثالث (القصوى 100)", bold=True, size=10, align=WD_ALIGN_PARAGRAPH.CENTER)
    format_paragraph(total_row_ax3.cells[4].paragraphs[0], f"{eval_res['axis3']['raw_total']}", bold=True, size=11, align=WD_ALIGN_PARAGRAPH.CENTER)

    format_paragraph(doc.add_paragraph(), "5 - 3", bold=True, size=9, align=WD_ALIGN_PARAGRAPH.CENTER)
    doc.add_page_break()

    # =========================================================================
    # الصفحة الرابعة : المحور الرابع (مواطن القوة)
    # =========================================================================
    p_ax4_head = doc.add_paragraph()
    format_paragraph(p_ax4_head, "المحور الرابع : مواطن القوة (تملى من قبل المسؤول المباشر) بعد ان تقدم الوثائق من قبل صاحب العلاقة المشمول بالتقييم\n(ملاحظة: في حال كون درجة الفقرة (المحور الثاني/1) صفر فان درجة التقييم النهائية يجب ان لا تتجاوز 75)", bold=True, size=10)

    ax4_table = doc.add_table(rows=14, cols=3)
    set_table_rtl(ax4_table)
    set_table_borders(ax4_table)
    ax4_table.alignment = WD_TABLE_ALIGNMENT.CENTER

    th_ax4 = ["ت", "مواطن القوة العلمية", "الدرجة المعطاة"]
    for i, h in enumerate(th_ax4):
        c = ax4_table.cell(0, i)
        set_cell_background(c, GRAY_HEX)
        format_paragraph(c.paragraphs[0], h, bold=True, size=10, align=WD_ALIGN_PARAGRAPH.CENTER)

    user_ax4_items = data.get("axis4", {}).get("items", {})
    strengths_definitions = [
        ("1", "براءات الاختراع و الجوائز (اي جائزة تم منحها ومطابقة في بياناتها لمتطلبات النظام الالكتروني) في عام التقييم حصرا (3 درجات)", user_ax4_items.get("item1", 0)),
        ("2", "امتلاك التدريسي لمعامل هيرش index -h على صفحته في Scopus (1: درجة، 2-3: درجتان، 4-6: 3 درجات، 7 فأكثر: 4 درجات)", user_ax4_items.get("item2", 0)),
        ("3", "مسؤول وحدة شؤون المرأة وجميع العاملين معهم (3 درجات)", user_ax4_items.get("item3", 0)),
        ("4", "تطوير منظومة الكترونية لإدارة احد البرامج على مستوى الجامعة او الوزارة (3 درجات)", user_ax4_items.get("item4", 0)),
        ("5", "مسؤولي الشعب والوحدات الارشادية واعضاء الارتباط (3 درجات)", user_ax4_items.get("item5", 0)),
        ("6", "اعضاء مجالس الاعتماد البرامجي والمؤسسي والمقيم الوطني واعضاء فريق تقييم الاداء المؤسسي (4 درجات)", user_ax4_items.get("item6", 0)),
        ("7", "مدراء اقسام وشعب ضمان الجودة والاداء الجامعي وجميع العاملين فيها (5 درجات)", user_ax4_items.get("item7", 0)),
        ("8", "المدرب المعتمد في طرائق التدريس من قبل وزارة التعليم العالي والبحث العلمي مشروطة بالتغذية الراجعة والتكليف الوزاري (5 درجات)", user_ax4_items.get("item8", 0)),
        ("9", "البحوث الاضافية المنشورة في المجلات المفهرسة في المستوعبات العالمية والمحلية والتي لم يتم الاستفادة منها في التقييم (4، 3، 2، 1 درجات)", user_ax4_items.get("item9", 0)),
        ("10", "دعم نادي الطلبة بمبلغ (2000) دينار شهريا (3 درجات)", user_ax4_items.get("item10", 0)),
        ("11", "الانتماء الى نقابة الاكاديميين (3 درجات)", user_ax4_items.get("item11", 0)),
        ("12", "مسؤولي شعب و وحدات حقوق الانسان (3 درجات)", user_ax4_items.get("item12", 0))
    ]

    for idx, (t, txt, scr) in enumerate(strengths_definitions, start=1):
        row = ax4_table.rows[idx]
        format_paragraph(row.cells[0].paragraphs[0], t, size=9, align=WD_ALIGN_PARAGRAPH.CENTER)
        format_paragraph(row.cells[1].paragraphs[0], txt, size=8)
        format_paragraph(row.cells[2].paragraphs[0], f"{scr}", bold=True, size=10, align=WD_ALIGN_PARAGRAPH.CENTER)

    total_row_ax4 = ax4_table.rows[13]
    total_row_ax4.cells[0].merge(total_row_ax4.cells[1])
    format_paragraph(total_row_ax4.cells[0].paragraphs[0], "الدرجة القصوى للمحور (5) درجات على ان لا تتجاوز درجة التقييم النهائية (100) درجة", bold=True, size=9, align=WD_ALIGN_PARAGRAPH.CENTER)
    format_paragraph(total_row_ax4.cells[2].paragraphs[0], f"{eval_res['axis4']['awarded_score']}", bold=True, size=11, align=WD_ALIGN_PARAGRAPH.CENTER)

    format_paragraph(doc.add_paragraph(), "5 - 4", bold=True, size=9, align=WD_ALIGN_PARAGRAPH.CENTER)
    doc.add_page_break()

    # =========================================================================
    # الصفحة الخامسة : المحور الخامس والنتائج والتواقيع
    # =========================================================================
    p_ax5_head = doc.add_paragraph()
    format_paragraph(p_ax5_head, "المحور الخامس : العقوبات (خصم الدرجات ) تملى من قبل المسؤول المباشر", bold=True, size=10)

    ax5_table = doc.add_table(rows=8, cols=3)
    set_table_rtl(ax5_table)
    set_table_borders(ax5_table)
    ax5_table.alignment = WD_TABLE_ALIGNMENT.CENTER

    th_ax5 = ["ت", "الإخفاق (تخصم الدرجة حسب الآتي)", "الدرجة التي تخصم"]
    for i, h in enumerate(th_ax5):
        c = ax5_table.cell(0, i)
        set_cell_background(c, GRAY_HEX)
        format_paragraph(c.paragraphs[0], h, bold=True, size=10, align=WD_ALIGN_PARAGRAPH.CENTER)

    penalties_data = [
        ("1", "لفت نظر", "تخصم (3) درجات"),
        ("2", "الإنذار", "تخصم (5) درجات"),
        ("3", "قطع الراتب", "تخصم (7) درجات"),
        ("4", "التوبيخ", "تخصم (11) درجة"),
        ("5", "إنقاص الراتب", "تخصم (13) درجة"),
        ("6", "تنزيل الدرجة", "تخصم (15) درجة")
    ]

    for idx, (t, name, ded_text) in enumerate(penalties_data, start=1):
        row = ax5_table.rows[idx]
        format_paragraph(row.cells[0].paragraphs[0], t, size=9, align=WD_ALIGN_PARAGRAPH.CENTER)
        format_paragraph(row.cells[1].paragraphs[0], f"{name} ({ded_text})", size=9)
        # البحث عن الخصم في المدخلات
        format_paragraph(row.cells[2].paragraphs[0], "-", size=9, align=WD_ALIGN_PARAGRAPH.CENTER)

    tot_pen_row = ax5_table.rows[7]
    tot_pen_row.cells[0].merge(tot_pen_row.cells[1])
    format_paragraph(tot_pen_row.cells[0].paragraphs[0], "المجموع الكلي للعقوبات المخصومة", bold=True, size=10, align=WD_ALIGN_PARAGRAPH.CENTER)
    format_paragraph(tot_pen_row.cells[2].paragraphs[0], f"-{eval_res['axis5']['total_deduction']}", bold=True, size=11, align=WD_ALIGN_PARAGRAPH.CENTER)

    p_res_head = doc.add_paragraph()
    p_res_head.paragraph_format.space_before = Pt(6)
    format_paragraph(p_res_head, "النتائج النهائية للتقييم :", bold=True, size=11)

    # جدول النتائج النهائية
    res_table = doc.add_table(rows=7, cols=4)
    set_table_rtl(res_table)
    set_table_borders(res_table)
    res_table.alignment = WD_TABLE_ALIGNMENT.CENTER

    th_res = ["ت", "المحاور", "وزن المحور", "الدرجة حسب الوزن"]
    for i, h in enumerate(th_res):
        c = res_table.cell(0, i)
        set_cell_background(c, YELLOW_HEX)
        format_paragraph(c.paragraphs[0], h, bold=True, size=10, align=WD_ALIGN_PARAGRAPH.CENTER)

    results_rows = [
        ("1", "جودة التدريس والتعليم والالتزام الوظيفي", "%50", f"{eval_res['axis1']['weighted_score']}"),
        ("2", "النشاط العلمي والبحثي", "%30", f"{eval_res['axis2']['weighted_score']}"),
        ("3", "الجانب التربوي والارشادي والتعليم المستمر", "%20", f"{eval_res['axis3']['weighted_score']}"),
        ("4", "مواطن القوة", "تضاف حتى 5 درجات", f"+{eval_res['axis4']['awarded_score']}"),
        ("5", "مجموع المحاور الثلاثة", "%100", f"{eval_res['three_axes_total']}"),
        ("6", "خصم درجات العقوبات", "تخصم بالكامل بدون وزن", f"-{eval_res['axis5']['total_deduction']}")
    ]

    for idx, (t, ax_name, weight, score_val) in enumerate(results_rows, start=1):
        row = res_table.rows[idx]
        format_paragraph(row.cells[0].paragraphs[0], t, size=9, align=WD_ALIGN_PARAGRAPH.CENTER)
        format_paragraph(row.cells[1].paragraphs[0], ax_name, bold=True, size=9)
        format_paragraph(row.cells[2].paragraphs[0], weight, size=9, align=WD_ALIGN_PARAGRAPH.CENTER)
        format_paragraph(row.cells[3].paragraphs[0], score_val, bold=True, size=10, align=WD_ALIGN_PARAGRAPH.CENTER)

    # الدرجة النهائية رقماً وكتابة
    final_score_table = doc.add_table(rows=1, cols=2)
    set_table_rtl(final_score_table)
    set_table_borders(final_score_table, sz="8")
    final_score_table.alignment = WD_TABLE_ALIGNMENT.CENTER

    c_s_num = final_score_table.cell(0, 0)
    c_s_txt = final_score_table.cell(0, 1)
    set_cell_background(c_s_num, LIGHT_YELLOW)
    set_cell_background(c_s_txt, LIGHT_YELLOW)

    format_paragraph(c_s_num.paragraphs[0], f"مجموع الدرجة رقماً:  {eval_res['final_score']} %", bold=True, size=12, align=WD_ALIGN_PARAGRAPH.CENTER)
    format_paragraph(c_s_txt.paragraphs[0], f"مجموع الدرجة كتابة:  {eval_res['score_in_words']}", bold=True, size=11, align=WD_ALIGN_PARAGRAPH.CENTER)

    # التقدير النهائي
    rating_table = doc.add_table(rows=2, cols=4)
    set_table_rtl(rating_table)
    set_table_borders(rating_table)
    rating_table.alignment = WD_TABLE_ALIGNMENT.CENTER

    rating_headers = ["امتياز (90 فاكثر)", "جيد جدا (80 - 89)", "جيد (70 - 79)", "ضعيف (اقل من 70)"]
    for i, rh in enumerate(rating_headers):
        c = rating_table.cell(0, i)
        set_cell_background(c, GRAY_HEX)
        format_paragraph(c.paragraphs[0], rh, bold=True, size=9, align=WD_ALIGN_PARAGRAPH.CENTER)

    # تحديد علامة صح على التقدير المستحق
    r_val = eval_res["rating"]
    markers = [
        "✔ (مستحق)" if r_val == "امتياز" else "",
        "✔ (مستحق)" if r_val == "جيد جدا" else "",
        "✔ (مستحق)" if r_val == "جيد" else "",
        "✔ (مستحق)" if r_val == "ضعيف" else ""
    ]
    for i, m in enumerate(markers):
        c = rating_table.cell(1, i)
        if m:
            set_cell_background(c, YELLOW_HEX)
        format_paragraph(c.paragraphs[0], m, bold=True, size=11, align=WD_ALIGN_PARAGRAPH.CENTER)

    # التوقيعات الرسمية
    p_sig = doc.add_paragraph()
    p_sig.paragraph_format.space_before = Pt(10)
    sig_table = doc.add_table(rows=3, cols=2)
    set_table_rtl(sig_table)
    sig_table.alignment = WD_TABLE_ALIGNMENT.CENTER

    format_paragraph(sig_table.cell(0, 0).paragraphs[0], "التوقيع : ...................................", bold=True, size=10)
    format_paragraph(sig_table.cell(0, 1).paragraphs[0], "التوقيع : ...................................", bold=True, size=10)
    format_paragraph(sig_table.cell(1, 0).paragraphs[0], "اسم المسؤول المباشر : ......................", bold=True, size=10)
    format_paragraph(sig_table.cell(1, 1).paragraphs[0], "اسم المسؤول الأعلى : ......................", bold=True, size=10)
    format_paragraph(sig_table.cell(2, 0).paragraphs[0], "التاريخ :       /       / 2026", bold=True, size=10)
    format_paragraph(sig_table.cell(2, 1).paragraphs[0], "التاريخ :       /       / 2026", bold=True, size=10)

    p_approv = doc.add_paragraph()
    p_approv.paragraph_format.space_before = Pt(10)
    format_paragraph(p_approv, "مصادقة رئيس الجامعة\nالاسم والتوقيع والختم الرسمي: .....................................................", bold=True, size=11, align=WD_ALIGN_PARAGRAPH.CENTER)

    format_paragraph(doc.add_paragraph(), "5 - 5", bold=True, size=9, align=WD_ALIGN_PARAGRAPH.CENTER)

    p_dev_note = doc.add_paragraph()
    p_dev_note.paragraph_format.space_before = Pt(4)
    format_paragraph(p_dev_note, "ملاحظة: هذه المنصة قيد التطوير وبمبادرة شخصية من المهندس المعماري الدكتور أحمد لؤي أحمد", size=8, italic=True, align=WD_ALIGN_PARAGRAPH.CENTER)

    # =========================================================================
    # صفحة الملحق : جدول المرفقات والوثائق الثبوتية المستخرجة
    # =========================================================================
    if attachments:
        doc.add_page_break()
        p_att_title = doc.add_paragraph()
        format_paragraph(p_att_title, "ملحق رقم (1): جدول المرفقات والوثائق الثبوتية المعالجة بنظام OCR", bold=True, size=13, align=WD_ALIGN_PARAGRAPH.CENTER)
        p_att_sub = doc.add_paragraph()
        format_paragraph(p_att_sub, f"اسم التدريسي: {personal.get('first_name', '')} {personal.get('father_name', '')} {personal.get('last_name', '')}    |    الكلية: {personal.get('college', '')}    |    القسم: {personal.get('department', '')}", size=10, align=WD_ALIGN_PARAGRAPH.CENTER)

        att_table = doc.add_table(rows=len(attachments) + 1, cols=8)
        set_table_rtl(att_table)
        set_table_borders(att_table)
        att_table.alignment = WD_TABLE_ALIGNMENT.CENTER

        th_att = ["ت", "رمز الفهرسة", "نوع الوثيقة", "العدد / الإشارة", "التاريخ", "الموضوع / النشاط", "المحور والفقرة", "الدرجة"]
        for i, h in enumerate(th_att):
            c = att_table.cell(0, i)
            set_cell_background(c, YELLOW_HEX)
            format_paragraph(c.paragraphs[0], h, bold=True, size=9, align=WD_ALIGN_PARAGRAPH.CENTER)

        for idx, att in enumerate(attachments, start=1):
            row = att_table.rows[idx]
            ref_code = att.get("ref_code", f"REF-{idx:02d}")
            format_paragraph(row.cells[0].paragraphs[0], str(idx), size=9, align=WD_ALIGN_PARAGRAPH.CENTER)
            format_paragraph(row.cells[1].paragraphs[0], ref_code, bold=True, size=8, align=WD_ALIGN_PARAGRAPH.CENTER)
            format_paragraph(row.cells[2].paragraphs[0], att.get("type_arabic", att.get("doc_type", att.get("type", "وثيقة"))), size=8)
            format_paragraph(row.cells[3].paragraphs[0], att.get("document_number", att.get("doc_number", "-")), size=8, align=WD_ALIGN_PARAGRAPH.CENTER)
            format_paragraph(row.cells[4].paragraphs[0], att.get("date", "-"), size=8, align=WD_ALIGN_PARAGRAPH.CENTER)
            format_paragraph(row.cells[5].paragraphs[0], att.get("title", att.get("filename", "")), size=8)
            format_paragraph(row.cells[6].paragraphs[0], f"{att.get('axis_name', att.get('axis', ''))} - فقرة {att.get('suggested_paragraph', att.get('paragraph', ''))}", size=8)
            format_paragraph(row.cells[7].paragraphs[0], str(att.get("suggested_score", 0)), bold=True, size=9, align=WD_ALIGN_PARAGRAPH.CENTER)

        p_att_dev = doc.add_paragraph()
        p_att_dev.paragraph_format.space_before = Pt(6)
        format_paragraph(p_att_dev, "هذه المنصة قيد التطوير وبمبادرة شخصية من المهندس المعماري الدكتور أحمد لؤي أحمد", size=8, italic=True, align=WD_ALIGN_PARAGRAPH.CENTER)

    # حفظ المستند
    doc.save(output_path)
    return output_path


def create_form_39_docx(data: Dict[str, Any], output_path: str, attachments: List[Dict[str, Any]] = None) -> str:
    """
    إنشاء ملف Word رسمي كامل لاستمارة رقم (39): تقييم أداء أعضاء الهيئة التدريسية لمسار بولونيا 2025-2026
    مطابق تماماً للجداول الرسمية الست والخلفيات ومحاذاة RTL وتفقيط الدرجة وملحق الوثائق.
    """
    eval_res = calculate_evaluation_form39(data)
    personal = data.get("personal_info", {})
    attachments = attachments or []

    doc = Document()

    # ضبط هوامش الصفحة
    for section in doc.sections:
        section.top_margin = Inches(0.4)
        section.bottom_margin = Inches(0.4)
        section.left_margin = Inches(0.4)
        section.right_margin = Inches(0.4)

    # =========================================================================
    # الصفحة الأولى (6 - 1) : الترويسة والبيانات الرئيسة
    # =========================================================================
    header_table = doc.add_table(rows=1, cols=2)
    set_table_rtl(header_table)
    header_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    header_table.autofit = True

    # يمين: الترويسة الوزارية
    c_right = header_table.cell(0, 0)
    p_r = c_right.paragraphs[0]
    format_paragraph(p_r, "وزارة التعليم العالي والبحث العلمي\nجهاز الإشراف والتقويم العلمي\nدائرة ضمان الجودة والاعتماد الأكاديمي\nقسم تقويم الأداء المؤسسي", bold=True, size=10, align=WD_ALIGN_PARAGRAPH.CENTER)

    # يسار: رقم وترميز الاستمارة
    c_left = header_table.cell(0, 1)
    p_l = c_left.paragraphs[0]
    form_no = personal.get("form_no", "2025/....")
    form_code = personal.get("form_code", "AGY-BOL-....")
    format_paragraph(p_l, f"رقم الاستمارة: {form_no}\nترميز الاستمارة: {form_code}", bold=True, size=10, align=WD_ALIGN_PARAGRAPH.LEFT)

    # عنوان الاستمارة الإطاري
    title_box = doc.add_table(rows=1, cols=1)
    set_table_rtl(title_box)
    set_table_borders(title_box, sz="8")
    c_title = title_box.cell(0, 0)
    p_t = c_title.paragraphs[0]
    format_paragraph(p_t, "استمارة رقم (39): تقييم أداء أعضاء الهيئة التدريسية لمسار بولونيا للعام الدراسي 2025-2026", bold=True, size=12, align=WD_ALIGN_PARAGRAPH.CENTER)
    set_cell_background(c_title, LIGHT_YELLOW)

    p_space = doc.add_paragraph()
    p_space.paragraph_format.space_before = Pt(4)
    format_paragraph(p_space, f"البيانات الرئيسة :\nالجامعة: {personal.get('university', 'الجامعة التكنولوجية')}    |    الكلية : {personal.get('college', 'هندسة العمارة')}    |    القسم /الفرع : {personal.get('department', 'قسم هندسة التصميم المعماري')}", bold=True, size=10)

    # جدول الاسم الرباعي واللقب (5 أعمدة)
    tbl_name = doc.add_table(rows=2, cols=5)
    set_table_rtl(tbl_name)
    set_table_borders(tbl_name)
    headers_name = ["اللقب", "اسم جد الاب", "اسم الجد", "اسم الاب", "الاسم"]
    for i, h in enumerate(headers_name):
        c = tbl_name.cell(0, i)
        set_cell_background(c, YELLOW_HEX)
        format_paragraph(c.paragraphs[0], h, bold=True, size=9, align=WD_ALIGN_PARAGRAPH.CENTER)
    values_name = [
        personal.get("last_name", ""),
        personal.get("great_grandfather_name", ""),
        personal.get("grandfather_name", ""),
        personal.get("father_name", ""),
        personal.get("first_name", "")
    ]
    for i, v in enumerate(values_name):
        format_paragraph(tbl_name.cell(1, i).paragraphs[0], v, size=9, align=WD_ALIGN_PARAGRAPH.CENTER)

    doc.add_paragraph().paragraph_format.space_before = Pt(2)

    # جدول اسم الأم الثلاثي (3 أعمدة)
    tbl_mother = doc.add_table(rows=2, cols=3)
    set_table_rtl(tbl_mother)
    set_table_borders(tbl_mother)
    headers_mother = ["اسم جد الام", "اسم والد الام", "اسم الام"]
    for i, h in enumerate(headers_mother):
        c = tbl_mother.cell(0, i)
        set_cell_background(c, YELLOW_HEX)
        format_paragraph(c.paragraphs[0], h, bold=True, size=9, align=WD_ALIGN_PARAGRAPH.CENTER)
    values_mother = [
        personal.get("mother_grandfather_name", ""),
        personal.get("mother_father_name", ""),
        personal.get("mother_name", "")
    ]
    for i, v in enumerate(values_mother):
        format_paragraph(tbl_mother.cell(1, i).paragraphs[0], v, size=9, align=WD_ALIGN_PARAGRAPH.CENTER)

    doc.add_paragraph().paragraph_format.space_before = Pt(2)

    # جدول الهوية والجنسية (4 أعمدة ثم صف التاريخ)
    tbl_id = doc.add_table(rows=4, cols=4)
    set_table_rtl(tbl_id)
    set_table_borders(tbl_id)
    
    # صف الهوية
    tbl_id.cell(0, 0).merge(tbl_id.cell(0, 1))
    set_cell_background(tbl_id.cell(0, 0), YELLOW_HEX)
    format_paragraph(tbl_id.cell(0, 0).paragraphs[0], "رقم الجنسية او البطاقة الموحدة", bold=True, size=9, align=WD_ALIGN_PARAGRAPH.CENTER)
    
    set_cell_background(tbl_id.cell(0, 2), YELLOW_HEX)
    format_paragraph(tbl_id.cell(0, 2).paragraphs[0], "السجل", bold=True, size=9, align=WD_ALIGN_PARAGRAPH.CENTER)
    
    set_cell_background(tbl_id.cell(0, 3), YELLOW_HEX)
    format_paragraph(tbl_id.cell(0, 3).paragraphs[0], "الصحيفة", bold=True, size=9, align=WD_ALIGN_PARAGRAPH.CENTER)

    tbl_id.cell(1, 0).merge(tbl_id.cell(1, 1))
    format_paragraph(tbl_id.cell(1, 0).paragraphs[0], personal.get("national_id", ""), size=9, align=WD_ALIGN_PARAGRAPH.CENTER)
    format_paragraph(tbl_id.cell(1, 2).paragraphs[0], personal.get("registry_no", ""), size=9, align=WD_ALIGN_PARAGRAPH.CENTER)
    format_paragraph(tbl_id.cell(1, 3).paragraphs[0], personal.get("page_no", ""), size=9, align=WD_ALIGN_PARAGRAPH.CENTER)

    # صف تاريخ الإصدار
    headers_issue = ["سنة الإصدار", "شهر الإصدار", "يوم الإصدار", "القسم"]
    for i, h in enumerate(headers_issue):
        c = tbl_id.cell(2, i)
        set_cell_background(c, YELLOW_HEX)
        format_paragraph(c.paragraphs[0], h, bold=True, size=9, align=WD_ALIGN_PARAGRAPH.CENTER)
    values_issue = [
        personal.get("issue_year", ""),
        personal.get("issue_month", ""),
        personal.get("issue_day", ""),
        personal.get("department", "")
    ]
    for i, v in enumerate(values_issue):
        format_paragraph(tbl_id.cell(3, i).paragraphs[0], v, size=9, align=WD_ALIGN_PARAGRAPH.CENTER)

    doc.add_paragraph().paragraph_format.space_before = Pt(2)

    # جدول الشهادة وتاريخ الأمر والجهة المانحة
    tbl_degree = doc.add_table(rows=4, cols=4)
    set_table_rtl(tbl_degree)
    set_table_borders(tbl_degree)
    headers_deg1 = ["الشهادة", "رقم وتاريخ الامر الوزاري او الجامعي", "شهر منح الشهادة", "يوم منح الشهادة"]
    for i, h in enumerate(headers_deg1):
        c = tbl_degree.cell(0, i)
        set_cell_background(c, YELLOW_HEX)
        format_paragraph(c.paragraphs[0], h, bold=True, size=9, align=WD_ALIGN_PARAGRAPH.CENTER)
    values_deg1 = [
        personal.get("degree", "دكتوراه"),
        personal.get("order_no_and_date", ""),
        personal.get("degree_month", ""),
        personal.get("degree_day", "")
    ]
    for i, v in enumerate(values_deg1):
        format_paragraph(tbl_degree.cell(1, i).paragraphs[0], v, size=9, align=WD_ALIGN_PARAGRAPH.CENTER)

    headers_deg2 = ["البلد المانح", "الجامعة", "الكلية", "القسم"]
    for i, h in enumerate(headers_deg2):
        c = tbl_degree.cell(2, i)
        set_cell_background(c, YELLOW_HEX)
        format_paragraph(c.paragraphs[0], h, bold=True, size=9, align=WD_ALIGN_PARAGRAPH.CENTER)
    values_deg2 = [
        personal.get("granting_country", "العراق"),
        personal.get("granting_univ", ""),
        personal.get("granting_college", ""),
        personal.get("granting_dept", "")
    ]
    for i, v in enumerate(values_deg2):
        format_paragraph(tbl_degree.cell(3, i).paragraphs[0], v, size=9, align=WD_ALIGN_PARAGRAPH.CENTER)

    doc.add_paragraph().paragraph_format.space_before = Pt(2)

    # جدول الاختصاص (2 أعمدة)
    tbl_spec = doc.add_table(rows=2, cols=2)
    set_table_rtl(tbl_spec)
    set_table_borders(tbl_spec)
    set_cell_background(tbl_spec.cell(0, 0), YELLOW_HEX)
    format_paragraph(tbl_spec.cell(0, 0).paragraphs[0], "الاختصاص العام", bold=True, size=9, align=WD_ALIGN_PARAGRAPH.CENTER)
    set_cell_background(tbl_spec.cell(0, 1), YELLOW_HEX)
    format_paragraph(tbl_spec.cell(0, 1).paragraphs[0], "الاختصاص الدقيق", bold=True, size=9, align=WD_ALIGN_PARAGRAPH.CENTER)
    format_paragraph(tbl_spec.cell(1, 0).paragraphs[0], personal.get("general_specialty", ""), size=9, align=WD_ALIGN_PARAGRAPH.CENTER)
    format_paragraph(tbl_spec.cell(1, 1).paragraphs[0], personal.get("specific_specialty", ""), size=9, align=WD_ALIGN_PARAGRAPH.CENTER)

    doc.add_paragraph().paragraph_format.space_before = Pt(2)

    # جدول اللقب العلمي والجهة المانحة وتاريخه
    tbl_title = doc.add_table(rows=2, cols=5)
    set_table_rtl(tbl_title)
    set_table_borders(tbl_title)
    headers_title = ["اللقب العلمي", "الجهة المانحة", "السنة", "شهر", "اليوم"]
    for i, h in enumerate(headers_title):
        c = tbl_title.cell(0, i)
        set_cell_background(c, YELLOW_HEX)
        format_paragraph(c.paragraphs[0], h, bold=True, size=9, align=WD_ALIGN_PARAGRAPH.CENTER)
    values_title = [
        personal.get("academic_title", "أستاذ مساعد"),
        personal.get("title_granter", personal.get("university", "الجامعة التكنولوجية")),
        personal.get("title_year", "2020"),
        personal.get("title_month", "05"),
        personal.get("title_day", "14")
    ]
    for i, v in enumerate(values_title):
        format_paragraph(tbl_title.cell(1, i).paragraphs[0], v, size=9, align=WD_ALIGN_PARAGRAPH.CENTER)

    doc.add_paragraph().paragraph_format.space_before = Pt(2)

    # جدول الاتصال
    tbl_contact = doc.add_table(rows=2, cols=2)
    set_table_rtl(tbl_contact)
    set_table_borders(tbl_contact)
    set_cell_background(tbl_contact.cell(0, 0), YELLOW_HEX)
    format_paragraph(tbl_contact.cell(0, 0).paragraphs[0], "رقم الموبايل", bold=True, size=9, align=WD_ALIGN_PARAGRAPH.CENTER)
    set_cell_background(tbl_contact.cell(0, 1), YELLOW_HEX)
    format_paragraph(tbl_contact.cell(0, 1).paragraphs[0], "البريد الالكتروني", bold=True, size=9, align=WD_ALIGN_PARAGRAPH.CENTER)
    format_paragraph(tbl_contact.cell(1, 0).paragraphs[0], personal.get("phone", ""), size=9, align=WD_ALIGN_PARAGRAPH.CENTER)
    format_paragraph(tbl_contact.cell(1, 1).paragraphs[0], personal.get("email", ""), size=9, align=WD_ALIGN_PARAGRAPH.CENTER)

    p_note1 = doc.add_paragraph()
    p_note1.paragraph_format.space_before = Pt(4)
    format_paragraph(p_note1, "ملاحظة: يتحمل صاحب الاستمارة مسؤولية صحة المعلومات الواردة في الاستمارة.", bold=True, size=9)

    p_f1 = doc.add_paragraph()
    format_paragraph(p_f1, "6 - 1", bold=True, size=9, align=WD_ALIGN_PARAGRAPH.CENTER)

    # =========================================================================
    # الصفحة الثانية (6 - 2) : المحور الأول والمحور الثاني (1 و 2)
    # =========================================================================
    doc.add_page_break()

    p_ax1_h = doc.add_paragraph()
    format_paragraph(p_ax1_h, "المحور الأول : تقييم الطالب للمادة الدراسية وطريقة التدريس (10 درجة)", bold=True, size=11, align=WD_ALIGN_PARAGRAPH.CENTER)

    tbl_ax1 = doc.add_table(rows=4, cols=5)
    set_table_rtl(tbl_ax1)
    set_table_borders(tbl_ax1)
    th_ax1 = ["ت", "الفقرات", "التوصيف", "تفصيل الدرجة", "الدرجة المعطاة"]
    for i, h in enumerate(th_ax1):
        c = tbl_ax1.cell(0, i)
        set_cell_background(c, YELLOW_HEX)
        format_paragraph(c.paragraphs[0], h, bold=True, size=9, align=WD_ALIGN_PARAGRAPH.CENTER)

    ax1_desc1 = "• يقوم المحاضر بتقديم موجز عن عملية التدريس والتقييم وفق مسار بولونيا للطلبة في أول محاضرة.\n• يقوم المحاضر بتقديم وشرح دليل وصف المادة الدراسية في الأسبوع الأول للفصل الدراسي.\n• يلتزم المحاضر بطريقة التقييم التكويني من حيث عدد الامتحانات القصيرة والواجبات المنزلية والفعاليات العلمية ومواعيدها والمذكورة في دليل وصف المادة الدراسية.\n• يقوم المحاضر بتوضيح علاقة مفردات المنهاج الأسبوعي بمخرجات التعلم.\nالدرجة القصوى (4) درجات درجة واحدة لكل فقرة"
    format_paragraph(tbl_ax1.cell(1, 0).paragraphs[0], "1", size=9, align=WD_ALIGN_PARAGRAPH.CENTER)
    format_paragraph(tbl_ax1.cell(1, 1).paragraphs[0], "التدريس وفق مسار بولونيا", bold=True, size=9)
    format_paragraph(tbl_ax1.cell(1, 2).paragraphs[0], ax1_desc1, size=8)
    format_paragraph(tbl_ax1.cell(1, 3).paragraphs[0], "4", size=9, align=WD_ALIGN_PARAGRAPH.CENTER)
    format_paragraph(tbl_ax1.cell(1, 4).paragraphs[0], str(eval_res["axis1"]["items"]["p1_bologna"]), bold=True, size=9, align=WD_ALIGN_PARAGRAPH.CENTER)

    ax1_desc2 = "• طريقة التدريس والمحاضرة مثيرة للاهتمام وتحفز الطالب ليفهم المادة العلمية.\n• يتم تزويد الطالب بقائمة من المراجع المختلفة بالإضافة للمراجع الرئيسية.\n• يعطي المحاضر الوقت الكافي للأسئلة والأجوبة خلال المحاضرة.\n• يستخدم المحاضر التقنيات اللازمة وأدوات الصوت والفيديو لشرح محاضراته.\n• خلال المحاضرة يعامل المحاضر الطلاب باحترام.\n• تعكس أسئلة الامتحان محتويات المادة العلمية المرتبطة بمخرجات التعلم.\nالدرجة القصوى (6) درجات درجة واحدة لكل فقرة"
    format_paragraph(tbl_ax1.cell(2, 0).paragraphs[0], "2", size=9, align=WD_ALIGN_PARAGRAPH.CENTER)
    format_paragraph(tbl_ax1.cell(2, 1).paragraphs[0], "كفاءة الأستاذ وطريقة التدريس", bold=True, size=9)
    format_paragraph(tbl_ax1.cell(2, 2).paragraphs[0], ax1_desc2, size=8)
    format_paragraph(tbl_ax1.cell(2, 3).paragraphs[0], "6", size=9, align=WD_ALIGN_PARAGRAPH.CENTER)
    format_paragraph(tbl_ax1.cell(2, 4).paragraphs[0], str(eval_res["axis1"]["items"]["p2_efficiency"]), bold=True, size=9, align=WD_ALIGN_PARAGRAPH.CENTER)

    tbl_ax1.cell(3, 0).merge(tbl_ax1.cell(3, 2))
    set_cell_background(tbl_ax1.cell(3, 0), GRAY_HEX)
    format_paragraph(tbl_ax1.cell(3, 0).paragraphs[0], "الدرجة النهائية للمحور الأول", bold=True, size=9, align=WD_ALIGN_PARAGRAPH.CENTER)
    format_paragraph(tbl_ax1.cell(3, 3).paragraphs[0], "10", bold=True, size=9, align=WD_ALIGN_PARAGRAPH.CENTER)
    format_paragraph(tbl_ax1.cell(3, 4).paragraphs[0], str(eval_res["axis1"]["total_score"]), bold=True, size=10, align=WD_ALIGN_PARAGRAPH.CENTER)

    p_ax2_h = doc.add_paragraph()
    p_ax2_h.paragraph_format.space_before = Pt(4)
    format_paragraph(p_ax2_h, "المحور الثاني: تقييم حقيبة الاستاذ (60 درجة)", bold=True, size=11, align=WD_ALIGN_PARAGRAPH.CENTER)

    tbl_ax2_p12 = doc.add_table(rows=4, cols=5)
    set_table_rtl(tbl_ax2_p12)
    set_table_borders(tbl_ax2_p12)
    th_ax2 = ["#", "الأبعاد الفرعية وقيمها القصوى", "المؤشرات وقيمها القصوى", "المقاييس", "الدرجة"]
    for i, h in enumerate(th_ax2):
        c = tbl_ax2_p12.cell(0, i)
        set_cell_background(c, YELLOW_HEX)
        format_paragraph(c.paragraphs[0], h, bold=True, size=9, align=WD_ALIGN_PARAGRAPH.CENTER)

    format_paragraph(tbl_ax2_p12.cell(1, 0).paragraphs[0], "1", size=9, align=WD_ALIGN_PARAGRAPH.CENTER)
    format_paragraph(tbl_ax2_p12.cell(1, 1).paragraphs[0], "واجبات ومسؤوليات أعضاء هيئة التدريس (8 درجات)", bold=True, size=9)
    format_paragraph(tbl_ax2_p12.cell(1, 2).paragraphs[0], "الحمل التدريسي (نصاب عضو هيئة التدريس)\n(4 درجات)", size=8)
    format_paragraph(tbl_ax2_p12.cell(1, 3).paragraphs[0], "• درجة كاملة في حال اكتمال النصاب\n• (2 درجة) في حال عدم اكتمال النصاب\n(ملاحظة: التدريسي الراغب وغير القادر قسمه يكلف بمهام داخل القسم وتعتبر له ساعات)", size=8)
    load_pts = 4.0 if data.get("axis2_portfolio", {}).get("teaching_load_status", "full") in ["full", "partial_department_assigned"] else 2.0
    format_paragraph(tbl_ax2_p12.cell(1, 4).paragraphs[0], str(load_pts), bold=True, size=9, align=WD_ALIGN_PARAGRAPH.CENTER)

    format_paragraph(tbl_ax2_p12.cell(2, 0).paragraphs[0], "", size=9)
    format_paragraph(tbl_ax2_p12.cell(2, 1).paragraphs[0], "تابع واجبات ومسؤوليات أعضاء هيئة التدريس", size=8)
    format_paragraph(tbl_ax2_p12.cell(2, 2).paragraphs[0], "عضوية اللجان العلمية في القسم والكلية والجامعة والوزارة (4 درجات)", size=8)
    format_paragraph(tbl_ax2_p12.cell(2, 3).paragraphs[0], "يتم تجميع الدرجات:\n• 4 درجات لكل لجنة دائمية\n• 2 درجة لكل لجنة مؤقتة", size=8)
    comm_pts = min(4.0, eval_res["axis2"]["items"]["p1_responsibilities"] - load_pts)
    format_paragraph(tbl_ax2_p12.cell(2, 4).paragraphs[0], str(max(0.0, comm_pts)), bold=True, size=9, align=WD_ALIGN_PARAGRAPH.CENTER)

    format_paragraph(tbl_ax2_p12.cell(3, 0).paragraphs[0], "2", size=9, align=WD_ALIGN_PARAGRAPH.CENTER)
    format_paragraph(tbl_ax2_p12.cell(3, 1).paragraphs[0], "كفاءة التدريسي في تحقيق مهارات ومعارف وقيم معينة (5 درجات)", bold=True, size=9)
    format_paragraph(tbl_ax2_p12.cell(3, 2).paragraphs[0], "كتب الشكر والتقدير والجوائز الأكاديمية (5 درجات)", size=8)
    format_paragraph(tbl_ax2_p12.cell(3, 3).paragraphs[0], "• درجتان لكتاب شكر من القسم\n• ثلاث درجات لكتاب شكر من الكلية\n• أربع درجات لكتاب شكر من الجامعة\n• خمس درجات لكتاب شكر من الوزارة\n• خمس درجات لجائزة أكاديمية", size=8)
    format_paragraph(tbl_ax2_p12.cell(3, 4).paragraphs[0], str(eval_res["axis2"]["items"]["p2_skills"]), bold=True, size=9, align=WD_ALIGN_PARAGRAPH.CENTER)

    p_f2 = doc.add_paragraph()
    format_paragraph(p_f2, "6 - 2", bold=True, size=9, align=WD_ALIGN_PARAGRAPH.CENTER)

    # =========================================================================
    # الصفحة الثالثة (6 - 3) : المحور الثاني (3 و 4)
    # =========================================================================
    doc.add_page_break()

    tbl_ax2_p34 = doc.add_table(rows=5, cols=5)
    set_table_rtl(tbl_ax2_p34)
    set_table_borders(tbl_ax2_p34)
    for i, h in enumerate(th_ax2):
        c = tbl_ax2_p34.cell(0, i)
        set_cell_background(c, YELLOW_HEX)
        format_paragraph(c.paragraphs[0], h, bold=True, size=9, align=WD_ALIGN_PARAGRAPH.CENTER)

    format_paragraph(tbl_ax2_p34.cell(1, 0).paragraphs[0], "3", size=9, align=WD_ALIGN_PARAGRAPH.CENTER)
    format_paragraph(tbl_ax2_p34.cell(1, 1).paragraphs[0], "التدريب والتعليم المستمر (20 درجة)", bold=True, size=9)
    format_paragraph(tbl_ax2_p34.cell(1, 2).paragraphs[0], "عقد الدورات التدريبية وورش العمل والندوات العلمية (6 درجات)", size=8)
    format_paragraph(tbl_ax2_p34.cell(1, 3).paragraphs[0], "• 4 درجات لعقد وتقديم دورة تدريبية لا تقل عن 24 ساعة وعن أسبوع\n• درجتان لعقد وتقديم ورشة عمل\n• درجتان لتقديم ندوة أو سمينار", size=8)
    format_paragraph(tbl_ax2_p34.cell(1, 4).paragraphs[0], "محتسب", size=8, align=WD_ALIGN_PARAGRAPH.CENTER)

    format_paragraph(tbl_ax2_p34.cell(2, 0).paragraphs[0], "", size=9)
    format_paragraph(tbl_ax2_p34.cell(2, 1).paragraphs[0], "تابع التدريب والتعليم المستمر", size=8)
    format_paragraph(tbl_ax2_p34.cell(2, 2).paragraphs[0], "حضور الدورات التدريبية او المؤتمرات العلمية ذات العلاقة بالتخصص (4 درجات)", size=8)
    format_paragraph(tbl_ax2_p34.cell(2, 3).paragraphs[0], "• درجتان لحضور دورة تدريبية\n• درجتان لحضور مؤتمر علمي مع تقديم ورقة علمية\n• درجة واحدة لحضور مؤتمر علمي بدون تقديم ورقة", size=8)
    format_paragraph(tbl_ax2_p34.cell(2, 4).paragraphs[0], "محتسب", size=8, align=WD_ALIGN_PARAGRAPH.CENTER)

    format_paragraph(tbl_ax2_p34.cell(3, 0).paragraphs[0], "", size=9)
    format_paragraph(tbl_ax2_p34.cell(3, 1).paragraphs[0], "تابع التدريب والتعليم المستمر", size=8)
    format_paragraph(tbl_ax2_p34.cell(3, 2).paragraphs[0], "حضور الندوات العلمية – سمينار وورش العمل (10 درجات)", size=8)
    format_paragraph(tbl_ax2_p34.cell(3, 3).paragraphs[0], "درجة واحدة لحضور سمينار او ورشة عمل داخل الجامعة مع تقديم فيدباك", size=8)
    format_paragraph(tbl_ax2_p34.cell(3, 4).paragraphs[0], str(eval_res["axis2"]["items"]["p3_training"]), bold=True, size=9, align=WD_ALIGN_PARAGRAPH.CENTER)

    res_desc = "• 20 درجة لبراءة اختراع دولية مسجلة في WIPO\n• 5 درجات لبراءة اختراع محلية\n• 8-12 درجة لبحث منشور في مجلة مؤرشفة في كلاريفيت وسكوبس معاً بحسب معامِل التأثير (أقل من 1: 8، 1-5: 10، أكثر من 5: 12)\n• 12 درجة لكتاب سكوبس/كلاريفيت أول، 8 ثان، 6 ثالث\n• 7 درجات لفصل كتاب أول، 4 ثان\n• 10 درجات لكتاب بطلب رسمي من الجامعة\n• 10-7 درجات لبحث سكوبس فقط (Q1: 10, Q2: 9, Q3: 8, Q4: 7)\n• 5 درجات لبحث غير مؤرشف مع DOI وميتاداتا وانتساب\nملاحظات:\n- لا يؤخذ ببحث لا يكون انتساب جامعته هو الأول عدا الدكتوراه\n- تضاف درجة للباحث الأول والمراسل (+1 لكل منهما)\n- خصم درجتان (-2) لتسلسل 5-10، وخصم 50% لتسلسل أكثر من 10\n- أقصى حد لتجميع درجات هذه الفقرة هو (20)"
    format_paragraph(tbl_ax2_p34.cell(4, 0).paragraphs[0], "4", size=9, align=WD_ALIGN_PARAGRAPH.CENTER)
    format_paragraph(tbl_ax2_p34.cell(4, 1).paragraphs[0], "رؤية الجامعة في البحث والتطوير ودعم القطاعات العلمية والصناعية (20 درجة)", bold=True, size=9)
    format_paragraph(tbl_ax2_p34.cell(4, 2).paragraphs[0], "نشر الابحاث والكتب وبراءات الاختراع (20 درجة)", size=8)
    format_paragraph(tbl_ax2_p34.cell(4, 3).paragraphs[0], res_desc, size=7)
    format_paragraph(tbl_ax2_p34.cell(4, 4).paragraphs[0], f"خام: {eval_res['axis2']['items']['p4_research_raw']}\nبعد المعامل: {eval_res['axis2']['items']['p4_research']}", bold=True, size=8, align=WD_ALIGN_PARAGRAPH.CENTER)

    p_f3 = doc.add_paragraph()
    format_paragraph(p_f3, "6 - 3", bold=True, size=9, align=WD_ALIGN_PARAGRAPH.CENTER)

    # =========================================================================
    # الصفحة الرابعة (6 - 4) : تابع المحور الثاني والمحور الثالث
    # =========================================================================
    doc.add_page_break()

    # جدول توضيحي لمعامل اللقب العلمي
    tbl_mult = doc.add_table(rows=2, cols=4)
    set_table_rtl(tbl_mult)
    set_table_borders(tbl_mult)
    headers_mult = ["أستاذ (100%)", "أستاذ مساعد (140%)", "مدرس / محاضر (160%)", "مدرس مساعد (180%)"]
    for i, h in enumerate(headers_mult):
        c = tbl_mult.cell(0, i)
        set_cell_background(c, YELLOW_HEX)
        format_paragraph(c.paragraphs[0], h, bold=True, size=8, align=WD_ALIGN_PARAGRAPH.CENTER)
    values_mult = [
        "مضروباً بـ 1.0",
        "مضروباً بـ 1.4",
        "مضروباً بـ 1.6",
        "مضروباً بـ 1.8"
    ]
    for i, v in enumerate(values_mult):
        format_paragraph(tbl_mult.cell(1, i).paragraphs[0], v, size=8, align=WD_ALIGN_PARAGRAPH.CENTER)

    p_ax2_cont = doc.add_paragraph()
    p_ax2_cont.paragraph_format.space_before = Pt(4)
    format_paragraph(p_ax2_cont, "المحور الثاني: الفقرات 5 و 6 (النشاط الدولي ودعم الإيرادات)", bold=True, size=10)

    tbl_ax2_p56 = doc.add_table(rows=5, cols=5)
    set_table_rtl(tbl_ax2_p56)
    set_table_borders(tbl_ax2_p56)
    for i, h in enumerate(th_ax2):
        c = tbl_ax2_p56.cell(0, i)
        set_cell_background(c, YELLOW_HEX)
        format_paragraph(c.paragraphs[0], h, bold=True, size=9, align=WD_ALIGN_PARAGRAPH.CENTER)

    format_paragraph(tbl_ax2_p56.cell(1, 0).paragraphs[0], "5", size=9, align=WD_ALIGN_PARAGRAPH.CENTER)
    format_paragraph(tbl_ax2_p56.cell(1, 1).paragraphs[0], "النشاط الدولي (5 درجات)", bold=True, size=9)
    format_paragraph(tbl_ax2_p56.cell(1, 2).paragraphs[0], "عضوية المؤسسات العلمية الدولية (درجتان)", size=8)
    format_paragraph(tbl_ax2_p56.cell(1, 3).paragraphs[0], "• درجتان لعضوية Senior or Fellow Member\n• درجة واحدة لعضوية Member", size=8)
    format_paragraph(tbl_ax2_p56.cell(1, 4).paragraphs[0], "محتسب", size=8, align=WD_ALIGN_PARAGRAPH.CENTER)

    format_paragraph(tbl_ax2_p56.cell(2, 0).paragraphs[0], "", size=9)
    format_paragraph(tbl_ax2_p56.cell(2, 1).paragraphs[0], "تابع النشاط الدولي", size=8)
    format_paragraph(tbl_ax2_p56.cell(2, 2).paragraphs[0], "عضوية لجان تحرير المجالت المؤرشفة وتقييم بحوثها (4 درجات)", size=8)
    format_paragraph(tbl_ax2_p56.cell(2, 3).paragraphs[0], "• 4 درجات لتحرير مجلة بكلاريفيت، 3 بسكوبس، 1 غير مؤرشفة\n• درجة واحدة لتقييم بحث لمجلة بسكوبس أو كلاريفيت", size=8)
    format_paragraph(tbl_ax2_p56.cell(2, 4).paragraphs[0], "محتسب", size=8, align=WD_ALIGN_PARAGRAPH.CENTER)

    format_paragraph(tbl_ax2_p56.cell(3, 0).paragraphs[0], "", size=9)
    format_paragraph(tbl_ax2_p56.cell(3, 1).paragraphs[0], "تابع النشاط الدولي", size=8)
    format_paragraph(tbl_ax2_p56.cell(3, 2).paragraphs[0], "متحدث رئيسي في المؤتمرات الدولية (درجتان)", size=8)
    format_paragraph(tbl_ax2_p56.cell(3, 3).paragraphs[0], "• درجتان كمتحدث رئيسي Keynote speaker بمؤتمر دولي\n• درجة واحدة بمؤتمر وطني", size=8)
    format_paragraph(tbl_ax2_p56.cell(3, 4).paragraphs[0], str(eval_res["axis2"]["items"]["p5_intl"]), bold=True, size=9, align=WD_ALIGN_PARAGRAPH.CENTER)

    format_paragraph(tbl_ax2_p56.cell(4, 0).paragraphs[0], "6", size=9, align=WD_ALIGN_PARAGRAPH.CENTER)
    format_paragraph(tbl_ax2_p56.cell(4, 1).paragraphs[0], "دعم إيرادات الجامعة (درجتان)", bold=True, size=9)
    format_paragraph(tbl_ax2_p56.cell(4, 2).paragraphs[0], "جلب الأموال إلى الجامعة (درجتان)", size=8)
    format_paragraph(tbl_ax2_p56.cell(4, 3).paragraphs[0], "درجة كاملة في حال الحصول على دعم مالي (Fund) للجامعة من خلال عمل أو مشروع", size=8)
    format_paragraph(tbl_ax2_p56.cell(4, 4).paragraphs[0], str(eval_res["axis2"]["items"]["p6_revenue"]), bold=True, size=9, align=WD_ALIGN_PARAGRAPH.CENTER)

    p_ax3_h = doc.add_paragraph()
    p_ax3_h.paragraph_format.space_before = Pt(4)
    format_paragraph(p_ax3_h, "3. المحور الثالث: تقييم رئيس القسم للعملية التعليمية واالستاذ الجامعي (15 درجة)", bold=True, size=11, align=WD_ALIGN_PARAGRAPH.CENTER)

    tbl_ax3 = doc.add_table(rows=3, cols=5)
    set_table_rtl(tbl_ax3)
    set_table_borders(tbl_ax3)
    th_ax3 = ["ت", "الفقرات", "التوصيف", "تفصيل الدرجة", "الدرجة المعطاة"]
    for i, h in enumerate(th_ax3):
        c = tbl_ax3.cell(0, i)
        set_cell_background(c, YELLOW_HEX)
        format_paragraph(c.paragraphs[0], h, bold=True, size=9, align=WD_ALIGN_PARAGRAPH.CENTER)

    ax3_desc = "• يشمل التقييم رأي رئيس القسم بمدى جدية وتعاون التدريسي مع القسم والعمل على تحقيق أهدافه وخصوصا فيما يخص تطبيق مسار بولونيا\n• في حال وجود عقوبات إدارية تؤخذ بنظر الاعتبار وكما يلي:\n- لفت نظر: خصم (3) درجات\n- الإنذار: خصم (4) درجات\n- قطع الراتب: خصم (5) درجات\n- التوبيخ: خصم (6) درجات\n- إنقاص الراتب: خصم (7) درجات\n- تنزيل الدرجة: خصم (8) درجات\nالدرجة القصوى للفقرة (15) درجة"
    format_paragraph(tbl_ax3.cell(1, 0).paragraphs[0], "1", size=9, align=WD_ALIGN_PARAGRAPH.CENTER)
    format_paragraph(tbl_ax3.cell(1, 1).paragraphs[0], "الالتزام برسالة القسم ورؤيته", bold=True, size=9)
    format_paragraph(tbl_ax3.cell(1, 2).paragraphs[0], ax3_desc, size=8)
    format_paragraph(tbl_ax3.cell(1, 3).paragraphs[0], "15", size=9, align=WD_ALIGN_PARAGRAPH.CENTER)
    format_paragraph(tbl_ax3.cell(1, 4).paragraphs[0], str(eval_res["axis3"]["total_score"]), bold=True, size=10, align=WD_ALIGN_PARAGRAPH.CENTER)

    tbl_ax3.cell(2, 0).merge(tbl_ax3.cell(2, 2))
    set_cell_background(tbl_ax3.cell(2, 0), GRAY_HEX)
    format_paragraph(tbl_ax3.cell(2, 0).paragraphs[0], "الدرجة النهائية للمحور الثالث", bold=True, size=9, align=WD_ALIGN_PARAGRAPH.CENTER)
    format_paragraph(tbl_ax3.cell(2, 3).paragraphs[0], "15", bold=True, size=9, align=WD_ALIGN_PARAGRAPH.CENTER)
    format_paragraph(tbl_ax3.cell(2, 4).paragraphs[0], str(eval_res["axis3"]["total_score"]), bold=True, size=10, align=WD_ALIGN_PARAGRAPH.CENTER)

    p_f4 = doc.add_paragraph()
    format_paragraph(p_f4, "6 - 4", bold=True, size=9, align=WD_ALIGN_PARAGRAPH.CENTER)

    # =========================================================================
    # الصفحة الخامسة (6 - 5) : المحور الرابع والنتائج النهائية والمصادقات
    # =========================================================================
    doc.add_page_break()

    p_ax4_h = doc.add_paragraph()
    format_paragraph(p_ax4_h, "4. المحور الرابع: الجانب التربوي والارشادي (15 درجة)", bold=True, size=11, align=WD_ALIGN_PARAGRAPH.CENTER)

    tbl_ax4 = doc.add_table(rows=4, cols=5)
    set_table_rtl(tbl_ax4)
    set_table_borders(tbl_ax4)
    for i, h in enumerate(th_ax1):
        c = tbl_ax4.cell(0, i)
        set_cell_background(c, YELLOW_HEX)
        format_paragraph(c.paragraphs[0], h, bold=True, size=9, align=WD_ALIGN_PARAGRAPH.CENTER)

    ax4_desc1 = "تمنح (5) درجات للمشاركة بصفة محاضر في التعليم المستمر والجودة، تمنح (4) درجات للمشاركة بصفة حضور في التعليم المستمر والجودة (فقط حضورين)، تمنح (3) درجات للمشاركة في دورات طرائق التدريس الحديثة في التعليم المستمر.\nالدرجة القصوى لهذه الفقرة (5) درجات"
    format_paragraph(tbl_ax4.cell(1, 0).paragraphs[0], "1", size=9, align=WD_ALIGN_PARAGRAPH.CENTER)
    format_paragraph(tbl_ax4.cell(1, 1).paragraphs[0], "المشاركة في لجان التعليم المستمر والجودة", bold=True, size=9)
    format_paragraph(tbl_ax4.cell(1, 2).paragraphs[0], ax4_desc1, size=8)
    format_paragraph(tbl_ax4.cell(1, 3).paragraphs[0], "5", size=9, align=WD_ALIGN_PARAGRAPH.CENTER)
    format_paragraph(tbl_ax4.cell(1, 4).paragraphs[0], str(eval_res["axis4"]["items"]["p1_qa"]), bold=True, size=9, align=WD_ALIGN_PARAGRAPH.CENTER)

    ax4_desc2 = "- تمنح (10) درجات لكل زيارة ميدانية للإشراف على الطلبة في مجال التطبيقات العملية أو التحليلات المختبرية أو السريرية أو السفرات العلمية.\n- تمنح (5) درجات لكل عمل تطوعي داخل وزارة التعليم العالي موثق بكتاب رسمي.\n- تمنح (5) درجات لكل تدقيق لغوي أو علمي أو إقامة ندوة أو ورشة أو مقال أو زيارة دار أيتام أو خدمة مستشفيات أو منظمات مجتمع مدني خارج وزارة التعليم العالي.\nالدرجة القصوى للفقرة (10) درجات"
    format_paragraph(tbl_ax4.cell(2, 0).paragraphs[0], "4", size=9, align=WD_ALIGN_PARAGRAPH.CENTER)
    format_paragraph(tbl_ax4.cell(2, 1).paragraphs[0], "الزيارات الميدانية والحقلية واالعمال التطوعية داخل الجامعة والمؤسسات العلمية او المجتمع", bold=True, size=9)
    format_paragraph(tbl_ax4.cell(2, 2).paragraphs[0], ax4_desc2, size=8)
    format_paragraph(tbl_ax4.cell(2, 3).paragraphs[0], "10", size=9, align=WD_ALIGN_PARAGRAPH.CENTER)
    format_paragraph(tbl_ax4.cell(2, 4).paragraphs[0], str(eval_res["axis4"]["items"]["p2_field"]), bold=True, size=9, align=WD_ALIGN_PARAGRAPH.CENTER)

    tbl_ax4.cell(3, 0).merge(tbl_ax4.cell(3, 2))
    set_cell_background(tbl_ax4.cell(3, 0), GRAY_HEX)
    format_paragraph(tbl_ax4.cell(3, 0).paragraphs[0], "الدرجة النهائية للمحور الرابع", bold=True, size=9, align=WD_ALIGN_PARAGRAPH.CENTER)
    format_paragraph(tbl_ax4.cell(3, 3).paragraphs[0], "15", bold=True, size=9, align=WD_ALIGN_PARAGRAPH.CENTER)
    format_paragraph(tbl_ax4.cell(3, 4).paragraphs[0], str(eval_res["axis4"]["total_score"]), bold=True, size=10, align=WD_ALIGN_PARAGRAPH.CENTER)

    p_res_h = doc.add_paragraph()
    p_res_h.paragraph_format.space_before = Pt(4)
    format_paragraph(p_res_h, "النتائج النهائية للتقييم :", bold=True, size=11)

    tbl_results = doc.add_table(rows=8, cols=3)
    set_table_rtl(tbl_results)
    set_table_borders(tbl_results)
    th_res = ["ت", "المحاور", "وزن المحور", "الدرجة"]
    # 3 cols: ت والمحاور | وزن المحور | الدرجة
    tbl_results.rows[0].cells[0].paragraphs[0]
    set_cell_background(tbl_results.cell(0, 0), YELLOW_HEX)
    format_paragraph(tbl_results.cell(0, 0).paragraphs[0], "ت / المحاور", bold=True, size=9, align=WD_ALIGN_PARAGRAPH.CENTER)
    set_cell_background(tbl_results.cell(0, 1), YELLOW_HEX)
    format_paragraph(tbl_results.cell(0, 1).paragraphs[0], "وزن المحور", bold=True, size=9, align=WD_ALIGN_PARAGRAPH.CENTER)
    set_cell_background(tbl_results.cell(0, 2), YELLOW_HEX)
    format_paragraph(tbl_results.cell(0, 2).paragraphs[0], "الدرجة المستحصلة", bold=True, size=9, align=WD_ALIGN_PARAGRAPH.CENTER)

    axes_rows = [
        ("1. تقييم الطالب", "%10", str(eval_res["axis1"]["weighted_score"])),
        ("2. حقيبة الاستاذ", "%60", str(eval_res["axis2"]["weighted_score"])),
        ("3. تقييم رئيس القسم للعملية التعليمية واالستاذ الجامعي", "%15", str(eval_res["axis3"]["weighted_score"])),
        ("4. الجانب التربوي والارشادي", "%15", str(eval_res["axis4"]["weighted_score"])),
        ("مجموع المحاور", "%100", str(eval_res["final_score"]))
    ]
    for idx, (m_name, m_weight, m_score) in enumerate(axes_rows, start=1):
        format_paragraph(tbl_results.cell(idx, 0).paragraphs[0], m_name, bold=(idx == 5), size=9)
        format_paragraph(tbl_results.cell(idx, 1).paragraphs[0], m_weight, bold=(idx == 5), size=9, align=WD_ALIGN_PARAGRAPH.CENTER)
        format_paragraph(tbl_results.cell(idx, 2).paragraphs[0], m_score, bold=True, size=9, align=WD_ALIGN_PARAGRAPH.CENTER)
        if idx == 5:
            set_cell_background(tbl_results.cell(idx, 0), GRAY_HEX)
            set_cell_background(tbl_results.cell(idx, 1), GRAY_HEX)
            set_cell_background(tbl_results.cell(idx, 2), GRAY_HEX)

    tbl_results.cell(6, 0).merge(tbl_results.cell(6, 1))
    format_paragraph(tbl_results.cell(6, 0).paragraphs[0], "مجموع الدرجات النهائية للتقييم (رقماً وتفقيطاً)", bold=True, size=9)
    format_paragraph(tbl_results.cell(6, 2).paragraphs[0], f"{eval_res['final_score']} ({eval_res['score_in_words']})", bold=True, size=9, align=WD_ALIGN_PARAGRAPH.CENTER)

    tbl_results.cell(7, 0).merge(tbl_results.cell(7, 1))
    set_cell_background(tbl_results.cell(7, 0), YELLOW_HEX)
    format_paragraph(tbl_results.cell(7, 0).paragraphs[0], "التقدير النهائي للتقييم", bold=True, size=10)
    rating_str = eval_res["rating"]
    set_cell_background(tbl_results.cell(7, 2), LIGHT_YELLOW)
    format_paragraph(tbl_results.cell(7, 2).paragraphs[0], f"★ {rating_str} ★", bold=True, size=11, align=WD_ALIGN_PARAGRAPH.CENTER)

    # جدول التقديرات الوزارية المرجعية
    tbl_ratings_bar = doc.add_table(rows=1, cols=4)
    set_table_rtl(tbl_ratings_bar)
    set_table_borders(tbl_ratings_bar)
    r_items = [
        ("امتياز (90 فأكثر)", rating_str == "امتياز"),
        ("جيد جدا (80-89)", rating_str == "جيد جدا"),
        ("جيد (70-79)", rating_str == "جيد"),
        ("ضعيف (اقل من 70)", rating_str == "ضعيف")
    ]
    for i, (rt, is_sel) in enumerate(r_items):
        c = tbl_ratings_bar.cell(0, i)
        if is_sel:
            set_cell_background(c, YELLOW_HEX)
        format_paragraph(c.paragraphs[0], f"{'✓ ' if is_sel else ''}{rt}", bold=is_sel, size=8, align=WD_ALIGN_PARAGRAPH.CENTER)

    # جدول التواقيع والمصادقات
    tbl_sigs = doc.add_table(rows=3, cols=2)
    set_table_rtl(tbl_sigs)
    tbl_sigs.alignment = WD_TABLE_ALIGNMENT.CENTER
    format_paragraph(tbl_sigs.cell(0, 0).paragraphs[0], "التوقيع : ...................................\nاسم المسؤول المباشر : .....................\nالتاريخ :        /      /  2026", size=8)
    format_paragraph(tbl_sigs.cell(0, 1).paragraphs[0], "التوقيع : ...................................\nاسم المسؤول الاعلى : ......................\nالتاريخ :        /      /  2026", size=8)
    
    tbl_sigs.cell(1, 0).merge(tbl_sigs.cell(1, 1))
    format_paragraph(tbl_sigs.cell(1, 0).paragraphs[0], "مصادقة رئيس الجامعة\n\n....................................................................", bold=True, size=9, align=WD_ALIGN_PARAGRAPH.CENTER)

    p_f5 = doc.add_paragraph()
    format_paragraph(p_f5, "6 - 5", bold=True, size=9, align=WD_ALIGN_PARAGRAPH.CENTER)

    # =========================================================================
    # الصفحة السادسة (6 - 6) : تعليمات وضوابط مسار بولونيا
    # =========================================================================
    doc.add_page_break()

    p_inst_h = doc.add_paragraph()
    format_paragraph(p_inst_h, "تعليمات وضوابط :", bold=True, size=12, align=WD_ALIGN_PARAGRAPH.RIGHT, color=RGBColor(180, 0, 0))

    instructions = [
        "1. تشمل الاستمارة تدريسيي مسار بولونيا من غير الإداريين والمجازين دراسياً أو وفق قانون 2020.",
        "2. حساب مؤشر الأداء الأكاديمي يكون مشروطاً بقيام عضو هيئة التدريس التسجيل بمنصة إدارة معلومات مسار بولونيا والتي تدار من قبل وزارة التعليم العالي والبحث العلمي وتقديم كافة المعلومات المطلوبة وبخلافه يحسب مؤشر الأداء الأكاديمي صفراً.",
        "3. أي عضو هيئة تدريس يقوم بتدريس مادة دراسية واحدة على الأقل يكون تحت طائلة حساب مؤشر الأداء الأكاديمي ومن خلال منصة إلكترونية تقوم الجامعة بتهيئتها لضمان الشفافية والكفاءة وتقديم كتب إثباتات النشاطات العلمية إلكترونياً.",
        "4. يتم أخذ تقييم الطالب حول المادة الدراسية وطريقة التدريس إلكترونياً وفي نهاية كل فصل دراسي وقبل بدء امتحانات نهاية الفصل (في الأسبوع العاشر) على أن تستمر عملية جمع تقييمات الطلبة لمدة أسبوع واحد.",
        "5. يتم حساب النصاب لعضو هيئة التدريس وفق عدد ساعات التدريس الفعلية وبتكليف من القسم بغض النظر عن تخصص المادة الدراسية وطبيعتها، وتراعى التعليمات الصادرة بخصوص قيمة الساعات النظرية مقابل الساعات العملية والمختبرية إن كان هنالك اختلاف.",
        "6. أي نشاط داخل القسم يتم تقاضي أجر لقائه لا يعتبر لغرض حساب مؤشر الأداء الأكاديمي، مثل الإشراف على رسائل الماجستير والدكتوراه ومناقشاتها وكذلك تقييم بحوث الترقية لقاء مبلغ مالي.",
        "7. لا يعتد بشهادات حضور السمينارات وحضور المؤتمرات العلمية او الدورات التدريبية التي تقام عبر الانترنت (أونلاين) أو خارج الجامعة عدا التي تقام بإشراف الوزارة او بموافقة الجامعة وبشكل رسمي.",
        "8. لا يتم احتساب حضور السمينارات داخل الجامعة بدون قيام التدريسيين الذين حضروا السمينار من تقديم فيدباك حول رأيهم بموضوع السمينار وكفاءة مقدم السمينار وطريقة العرض والجوانب الأخرى وخلال فترة لا تتجاوز الأسبوع من تاريخ تقديم السمينار وذلك من خلال استمارة إلكترونية يتم إعدادها للغرض من قبل ضمان جودة الجامعة أو الكلية.",
        "9. على رئيس القسم الاطلاع دورياً على حسابات نظام إدارة التعلم الخاص بالتدريسيين في قسمه وبالتعاون مع مدير النظام.",
        "10. بخصوص الفقرة الرابعة 'رؤية الجامعة في البحث والتطوير ودعم القطاعات العلمية والصناعية' والواردة في تقييم حقيبة الأستاذ، يتم ضرب درجة هذه الفقرة حسب الآتي:\n   - مضروباً بـ 100% للتدريسي بدرجة 'أستاذ'\n   - مضروباً بـ 140% للتدريسي بدرجة 'أستاذ مساعد'\n   - مضروباً بـ 160% للتدريسي بدرجة 'مدرس' (محاضر)\n   - مضروباً بـ 180% للتدريسي بدرجة 'مدرس مساعد' (محاضر مساعد)",
        "11. تلتزم الجامعة بتعليمات الوزارة الصادرة بخصوص النشر العلمي والنشر في المجلات التجارية أو المجلات المفترسة أو المجلات التي تدرج ضمن القائمة السوداء.",
        "12. تكون هذه الوثيقة جزء من 'سياسة ضمان جودة مسار بولونيا' والتي يجب على الجامعة أو المؤسسة التعليمية كتابتها وإصدارها وجعلها متاحة لجميع أعضاء هيئة التدريس على موقع الجامعة وذلك حسب متطلبات المعيار الأول من تقرير ضمان وتعزيز الجودة السنوي."
    ]

    for inst in instructions:
        p_in = doc.add_paragraph()
        p_in.paragraph_format.space_before = Pt(2)
        p_in.paragraph_format.space_after = Pt(2)
        format_paragraph(p_in, inst, size=8.5)

    p_f6 = doc.add_paragraph()
    format_paragraph(p_f6, "6 - 6", bold=True, size=9, align=WD_ALIGN_PARAGRAPH.CENTER)

    # =========================================================================
    # صفحة الملحق التوثيقي : جدول المرفقات والوثائق الثبوتية المعالجة بـ OCR
    # =========================================================================
    if attachments:
        doc.add_page_break()
        p_att_title = doc.add_paragraph()
        format_paragraph(p_att_title, "ملحق التوثيق والأدلة: جدول الوثائق والمرفقات المستخرجة بنظام OCR لمسار بولونيا", bold=True, size=12, align=WD_ALIGN_PARAGRAPH.CENTER)
        p_att_sub = doc.add_paragraph()
        format_paragraph(p_att_sub, f"اسم التدريسي: {personal.get('first_name', '')} {personal.get('father_name', '')} {personal.get('last_name', '')}    |    الكلية: {personal.get('college', '')}    |    القسم: {personal.get('department', '')}", size=9, align=WD_ALIGN_PARAGRAPH.CENTER)

        att_table = doc.add_table(rows=len(attachments) + 1, cols=8)
        set_table_rtl(att_table)
        set_table_borders(att_table)
        att_table.alignment = WD_TABLE_ALIGNMENT.CENTER

        th_att = ["ت", "رمز الفهرسة", "نوع الوثيقة", "العدد / الإشارة", "التاريخ", "الموضوع / النشاط", "المحور والفقرة", "الدرجة"]
        for i, h in enumerate(th_att):
            c = att_table.cell(0, i)
            set_cell_background(c, YELLOW_HEX)
            format_paragraph(c.paragraphs[0], h, bold=True, size=8.5, align=WD_ALIGN_PARAGRAPH.CENTER)

        for idx, att in enumerate(attachments, start=1):
            row = att_table.rows[idx]
            ref_code = att.get("ref_code", f"REF-P39-{idx:02d}")
            format_paragraph(row.cells[0].paragraphs[0], str(idx), size=8, align=WD_ALIGN_PARAGRAPH.CENTER)
            format_paragraph(row.cells[1].paragraphs[0], ref_code, bold=True, size=8, align=WD_ALIGN_PARAGRAPH.CENTER)
            format_paragraph(row.cells[2].paragraphs[0], att.get("type_arabic", att.get("doc_type", att.get("type", "وثيقة"))), size=8)
            format_paragraph(row.cells[3].paragraphs[0], att.get("document_number", att.get("doc_number", "-")), size=8, align=WD_ALIGN_PARAGRAPH.CENTER)
            format_paragraph(row.cells[4].paragraphs[0], att.get("date", "-"), size=8, align=WD_ALIGN_PARAGRAPH.CENTER)
            format_paragraph(row.cells[5].paragraphs[0], att.get("title", att.get("filename", "")), size=8)
            format_paragraph(row.cells[6].paragraphs[0], f"{att.get('axis_name', att.get('axis', ''))} - فقرة {att.get('suggested_paragraph', att.get('paragraph', ''))}", size=8)
            format_paragraph(row.cells[7].paragraphs[0], str(att.get("suggested_score", 0)), bold=True, size=8, align=WD_ALIGN_PARAGRAPH.CENTER)

    # حفظ المستند
    doc.save(output_path)
    return output_path

